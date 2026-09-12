#!/usr/bin/env python3
"""把已有 DOCX 迁移到本技能的样式基线并输出新文件。

脚本只在内存中读取输入，绝不覆盖输入文件；正文文字、表格、图片和 OMML
对象保留，手工标题编号作为非内容格式被移除后由样式编号重新生成。
"""

from __future__ import annotations

import argparse
import copy
import re
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from build_base_template import (
    add_baseline_marker,
    clear_container,
    configure_document,
    configure_styles,
    find,
    remove_all,
    tag,
)


M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
BASE_TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "assets" / "chinese-docx-base-template.docx"


def clean(text: str | None) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def style_text(paragraph) -> str:
    style = paragraph.style
    return f"{style.style_id} {style.name}".lower() if style is not None else ""


def in_table(paragraph) -> bool:
    node = paragraph._p
    while node is not None:
        if node.tag == tag("tc"):
            return True
        node = node.getparent()
    return False


def is_layout_only_paragraph(paragraph) -> bool:
    """识别只含空格、制表位或分页符的正文段落。"""
    if (paragraph.text or "").strip(" \t\r\n\u00a0"):
        return False
    meaningful = {
        tag_name
        for tag_name in (
            tag("drawing"),
            tag("object"),
            tag("pict"),
            tag("oMath"),
            tag("oMathPara"),
            tag("instrText"),
            tag("fldSimple"),
            tag("hyperlink"),
            tag("footnoteReference"),
            tag("endnoteReference"),
        )
    }
    return not any(node.tag in meaningful for node in paragraph._p.iter())


def meaningful_paragraphs(document: Document) -> list:
    """过滤正文中的版式占位段落，但保留表格单元格内的空段落。"""
    return [
        paragraph
        for paragraph in all_paragraphs(document)
        if in_table(paragraph) or not is_layout_only_paragraph(paragraph)
    ]


def has_omml(paragraph) -> bool:
    return any(node.tag.startswith("{" + M + "}") for node in paragraph._p.iter())


def is_toc_entry(paragraph) -> bool:
    """只识别目录样式或带页码制表位的目录项，避免误伤引用超链接。"""
    style = style_text(paragraph)
    if any(marker in style for marker in ("toc", "table of figures", "table of tables", "目录")):
        return True
    return any(node.tag == tag("hyperlink") for node in paragraph._p.iter()) and any(
        node.tag == tag("tab") for node in paragraph._p.iter()
    )


def toc_role(paragraph) -> str | None:
    if not is_toc_entry(paragraph):
        return None
    style = style_text(paragraph)
    if "table of figures" in style:
        return "Figure List"
    if "table of tables" in style:
        return "Table List"
    text = clean(paragraph.text)
    if re.match(r"^\d+\.\d+\.\d+", text):
        return "TOC 3"
    if re.match(r"^\d+\.\d+", text):
        return "TOC 2"
    return "TOC 1"


def all_table_paragraphs(tables, result: list) -> None:
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph not in result:
                        result.append(paragraph)
                all_table_paragraphs(cell.tables, result)


def all_paragraphs(document: Document) -> list:
    result = list(document.paragraphs)
    all_table_paragraphs(document.tables, result)
    unique = []
    seen = set()
    for paragraph in result:
        key = id(paragraph._p)
        if key not in seen:
            seen.add(key)
            unique.append(paragraph)
    return unique


def detect_profile(document: Document) -> str:
    """按论文结构信号选择模式，不依赖源文档的旧样式名称。"""
    texts = [clean(paragraph.text) for paragraph in all_paragraphs(document)]
    joined = "\n".join(texts)
    signals = (
        bool(re.search(r"(?:摘\s*要|摘要)", joined)),
        "ABSTRACT" in joined.upper(),
        bool(re.search(r"(?:目\s*录|目录)", joined)),
        "参考文献" in joined,
        any(re.match(r"^第[一二三四五六七八九十百千万零〇]+章", text) for text in texts),
    )
    return "thesis" if sum(signals) >= 2 else "general"


def clear_main_body(document: Document) -> None:
    """清空基线模板正文，同时保留唯一节属性。"""
    body = document._body._element
    for child in list(body):
        if child.tag != tag("sectPr"):
            body.remove(child)


def source_body_content(source: Document) -> list:
    """复制正文和表格节点，不复制源文档的节属性。"""
    content = []
    top_level_paragraphs = {id(paragraph._p): paragraph for paragraph in source.paragraphs}
    for child in list(source._body._element):
        if child.tag == tag("sectPr"):
            continue
        if child.tag == tag("p"):
            paragraph = top_level_paragraphs.get(id(child))
            if paragraph is not None and is_layout_only_paragraph(paragraph):
                continue
        node = copy.deepcopy(child)
        for marker in list(node.iter(tag("lastRenderedPageBreak"))):
            parent = marker.getparent()
            if parent is not None:
                parent.remove(marker)
        for sect_pr in list(node.iter(tag("sectPr"))):
            parent = sect_pr.getparent()
            if parent is not None:
                parent.remove(sect_pr)
        content.append(node)
    return content


def relationship_ids(root) -> set[str]:
    return {
        value
        for node in root.iter()
        for name, value in node.attrib.items()
        if name.startswith("{" + REL_NS + "}")
    }


def clone_related_part(source_part, target_package, cache: dict):
    """将正文引用的媒体或嵌入对象复制到目标包。"""
    key = str(source_part.partname)
    if key in cache:
        return cache[key]

    existing = next(
        (part for part in target_package.iter_parts() if str(part.partname) == key),
        None,
    )
    if existing is not None:
        cache[key] = existing
        return existing

    cloned = Part(
        source_part.partname,
        source_part.content_type,
        source_part.blob,
        target_package,
    )
    cache[key] = cloned
    return cloned


def copy_source_relationships(source: Document, target: Document, content: list) -> dict[str, str]:
    """复制正文所需关系，并返回源 rId 到目标 rId 的映射。"""
    wrapper = OxmlElement("w:body")
    for node in content:
        wrapper.append(copy.deepcopy(node))
    needed = relationship_ids(wrapper)
    source_body = source._body._element
    has_footnotes = any(node.tag == tag("footnoteReference") for node in source_body.iter())
    has_endnotes = any(node.tag == tag("endnoteReference") for node in source_body.iter())
    for rel_id, relation in source.part.rels.items():
        if has_footnotes and relation.reltype.endswith("/footnotes"):
            needed.add(rel_id)
        if has_endnotes and relation.reltype.endswith("/endnotes"):
            needed.add(rel_id)

    target_part = target.part
    target_package = target.part.package
    part_cache = {}
    mapping = {}
    for rel_id in sorted(needed):
        relation = source.part.rels.get(rel_id)
        if relation is None:
            continue
        if relation.is_external:
            new_id = target_part.relate_to(
                relation.target_ref,
                relation.reltype,
                is_external=True,
            )
        else:
            cloned = clone_related_part(relation.target_part, target_package, part_cache)
            new_id = target_part.relate_to(cloned, relation.reltype)
        mapping[rel_id] = new_id
    return mapping


def remap_relationship_ids(root, mapping: dict[str, str]) -> None:
    for node in root.iter():
        for name, value in list(node.attrib.items()):
            if name.startswith("{" + REL_NS + "}") and value in mapping:
                node.set(name, mapping[value])


def first_nonempty_header_text(document: Document, attribute: str) -> str:
    for section in document.sections:
        container = getattr(section, attribute)
        value = clean(" ".join(paragraph.text for paragraph in container.paragraphs))
        if value:
            return value
    return ""


def set_header_text(container, text: str, style_id: str) -> None:
    if not text:
        return
    clear_container(container, style_id)
    container.paragraphs[0].text = text


def copy_thesis_headers(source: Document, target: Document) -> None:
    """论文模式只保留源文档普通页眉文字，不创建特殊页眉引用。"""
    section = target.sections[0]
    header_id = target.styles["Header"].style_id
    set_header_text(section.header, first_nonempty_header_text(source, "header"), header_id)


def looks_like_heading(text: str) -> int | None:
    if not text or len(text) > 100 or text.endswith(("。", "；", ";", "！", "!", "？", "?")):
        return None
    if re.match(r"^第[一二三四五六七八九十百千万零〇]+章(?:\s|$)", text):
        return 1
    if re.match(r"^[一二三四五六七八九十百千万零〇]+[、.．]\s*", text):
        return 1
    if re.match(r"^\d+\.\d+\.\d+(?:\s|$)", text):
        return 3
    if re.match(r"^\d+\.\d+(?:\s|$)", text):
        return 2
    return None


def is_technical(text: str, raw_text: str | None = None) -> bool:
    if not text:
        return False
    raw_text = raw_text or text
    normalized_raw = raw_text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.strip() for line in normalized_raw.split("\n") if line.strip()]
    technical_markers = (
        "::",
        "=>",
        "->",
        "↓",
        "{",
        "}",
        "GET ",
        "POST ",
        "PUT ",
        "DELETE ",
        "SELECT ",
        "import ",
        "from ",
        "def ",
        "class ",
        "curl ",
        "python ",
        "mvn ",
        "npm ",
    )
    line_markers = (
        "GET ",
        "POST ",
        "PUT ",
        "DELETE ",
        "SELECT ",
        "import ",
        "from ",
        "def ",
        "class ",
        "curl ",
        "python ",
        "mvn ",
        "npm ",
    )
    if any(line.startswith(line_markers) for line in lines):
        return True
    marker_count = sum(text.count(marker) for marker in technical_markers)
    path_count = text.count("/") + text.count("\\")
    if any(line.startswith(("├", "└", "│", "↓", "↳")) for line in lines):
        return True
    if "\n" in normalized_raw and (marker_count >= 1 or path_count >= 2):
        return True
    if marker_count >= 2:
        return True
    if len(lines) > 1 and path_count >= 3:
        return True
    return False


def classify(paragraph, is_first_nonempty: bool) -> str:
    text = clean(paragraph.text)
    raw_text = paragraph.text or ""
    style = style_text(paragraph)
    toc = toc_role(paragraph)
    if toc:
        return toc
    if has_omml(paragraph):
        return "Equation"
    if any(marker in style for marker in ("code block", "codeblock", "codepath", "代码", "源码")):
        return "Code Block"
    if "figure caption" in style or "图题" in style or re.match(r"^图\s*[0-9A-Za-z]", text):
        return "Figure Caption"
    if "table caption" in style or "表题" in style or re.match(r"^表\s*[0-9A-Za-z]", text):
        return "Table Caption"
    if "reference" in style or "endnote bibliography" in style:
        return "Reference"
    if "appendix heading" in style or re.match(r"^附录\s*[0-9A-Za-z]", text):
        return "Appendix Heading"
    if text in {"摘要", "摘 要"}:
        return "Abstract Heading"
    if text == "ABSTRACT":
        return "English Abstract Heading"
    compact = re.sub(r"\s+", "", text)
    if compact in {"目录", "图目录", "表目录", "插图目录", "插表目录"}:
        return "TOC Title"
    if text == "参考文献":
        return "Heading 1"
    if "abstract body" in style and "english" in style:
        return "English Abstract Body"
    if "abstract body" in style:
        return "Abstract Body"
    if "heading 1" in style or "heading1" in style or "标题 1" in style:
        return "Heading 1"
    if "heading 2" in style or "heading2" in style or "标题 2" in style:
        return "Heading 2"
    if "heading 3" in style or "heading3" in style or "标题 3" in style:
        return "Heading 3"
    if "heading 4" in style or "heading4" in style or "标题 4" in style:
        # 本技能只定义三级标题；将源文档第四级标题纳入三级样式，避免降为正文。
        return "Heading 3"
    if "title" in style and is_first_nonempty:
        return "Title"
    if is_first_nonempty and not in_table(paragraph) and text and len(text) <= 100:
        return "Title"
    if not in_table(paragraph) and looks_like_heading(text):
        return f"Heading {looks_like_heading(text)}"
    if is_technical(text, raw_text):
        return "Code Block"
    if in_table(paragraph):
        return "Table Text"
    return "Normal"


def is_front_matter_metadata(text: str) -> bool:
    """跳过论文封面上的编号、分类和日期字段，避免把它们当作文档标题。"""
    if not text or len(text) > 80:
        return False
    compact = re.sub(r"\s+", "", text).upper()
    markers = (
        "学校代码",
        "学号",
        "分类号",
        "密级",
        "UDC",
        "论文题目",
        "学位论文",
        "申请学位",
        "专业",
        "学科",
        "作者",
        "指导教师",
        "答辩日期",
        "提交日期",
    )
    return any(marker.upper() in compact for marker in markers)


def find_title_index(paragraphs: list, profile: str) -> int | None:
    for index, paragraph in enumerate(paragraphs):
        text = clean(paragraph.text)
        if not text:
            continue
        if "title" in style_text(paragraph):
            return index
        if profile == "thesis" and is_front_matter_metadata(text):
            continue
        return index
    return None


def classify_paragraphs(paragraphs: list, profile: str) -> list[tuple[object, str]]:
    """按段落上下文识别摘要正文，同时保留源段落的顺序。"""
    title_index = find_title_index(paragraphs, profile)
    module: str | None = None
    roles = []
    for index, paragraph in enumerate(paragraphs):
        text = clean(paragraph.text)
        compact = re.sub(r"\s+", "", text)
        role = classify(paragraph, index == title_index)
        if compact in {"摘要"}:
            module = "abstract"
            role = "Abstract Heading"
        elif compact == "ABSTRACT":
            module = "english"
            role = "English Abstract Heading"
        elif compact.startswith("关键词"):
            module = None
        elif compact.startswith("KEYWORDS"):
            module = None
        elif module == "abstract" and text and role == "Normal" and not in_table(paragraph):
            role = "Abstract Body"
        elif module == "english" and text and role == "Normal" and not in_table(paragraph):
            role = "English Abstract Body"
        roles.append((paragraph, role))
    return roles


def strip_heading_prefix(paragraph, role: str) -> None:
    if role not in {"Heading 1", "Heading 2", "Heading 3", "Appendix Heading"}:
        return
    text = paragraph.text or ""
    patterns = {
        "Heading 1": r"^\s*(?:第[一二三四五六七八九十百千万零〇]+章|[一二三四五六七八九十百千万零〇]+[、.．]|\d+\.)\s*",
        "Heading 2": r"^\s*\d+\.\d+\s*",
        "Heading 3": r"^\s*\d+\.\d+\.\d+\s*",
        "Appendix Heading": r"^\s*附录\s*[0-9A-Za-z]+\s*",
    }
    leading = re.match(r"^[ \t\u00a0]+", text)
    leading_text = leading.group(0) if leading else ""
    remainder = text[len(leading_text):]
    match = re.match(patterns[role], remainder)
    prefix = leading_text + (match.group(0) if match else "")
    if not prefix or not paragraph.runs:
        return
    remaining = len(prefix)
    for run in paragraph.runs:
        if remaining <= 0:
            break
        value = run.text or ""
        if not value:
            continue
        remove_count = min(remaining, len(value))
        run.text = value[remove_count:]
        remaining -= remove_count


def strip_paragraph_edge_whitespace(paragraph) -> None:
    """移除非代码段落首尾的布局空白，不破坏字段和 OMML。"""
    nodes = list(paragraph._p.iter())
    leading = True
    for node in nodes:
        if node.tag == tag("t"):
            value = node.text or ""
            if leading:
                value = value.lstrip(" \t\u00a0")
                node.text = value or None
                if value:
                    leading = False
        elif node.tag == tag("tab") and leading:
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)

    trailing = True
    for node in reversed(nodes):
        if node.getparent() is None:
            continue
        if node.tag == tag("t"):
            value = node.text or ""
            if trailing:
                value = value.rstrip(" \t\u00a0")
                node.text = value or None
                if value:
                    trailing = False
        elif node.tag == tag("tab") and trailing:
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)


def clean_paragraph_properties(paragraph, style_name: str, style_id: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    for child in list(ppr):
        if child.tag != tag("sectPr"):
            ppr.remove(child)
    ppr = paragraph._p.get_or_add_pPr()
    pstyle = OxmlElement("w:pStyle")
    pstyle.set(tag("val"), style_id)
    ppr.insert(0, pstyle)


def clean_run_properties(paragraph) -> None:
    for marker in list(paragraph._p.iter(tag("proofErr"))):
        parent = marker.getparent()
        if parent is not None:
            parent.remove(marker)
    for run in paragraph._p.iter(tag("r")):
        rpr = run.find(tag("rPr"))
        if rpr is None:
            continue
        # 保留上标、下标和语言信息，清除会覆盖样式的视觉属性。
        for child in list(rpr):
            if child.tag not in {tag("vertAlign"), tag("lang"), tag("rtl")}:
                rpr.remove(child)
        if len(rpr) == 0:
            run.remove(rpr)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = find(tc_pr, "tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.insert(0, tc_w)
    tc_w.set(tag("w"), str(width_twips))
    tc_w.set(tag("type"), "dxa")


def visual_text_width(text: str) -> float:
    """按中英文混排的近似字面宽度计算表格列需求。"""
    width = 0.0
    for char in text:
        if char.isspace():
            width += 0.35
        elif unicodedata.east_asian_width(char) in {"W", "F"}:
            width += 1.0
        else:
            width += 0.55
    return width


def table_column_score(table, column_index: int) -> float:
    values = []
    for row in table.rows:
        if column_index >= len(row.cells):
            continue
        values.extend(paragraph.text for paragraph in row.cells[column_index].paragraphs)
    if not values:
        return 1.0
    line_width = max((visual_text_width(line) for value in values for line in value.splitlines()), default=1.0)
    tokens = [token for value in values for token in re.split(r"[\s/\\,，;；:：]+", value) if token]
    longest_token = max((visual_text_width(token) for token in tokens), default=1.0)
    return max(1.0, min(60.0, line_width + longest_token * 0.35))


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = find(tbl_pr, "tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = find(borders, side)
        if node is None:
            node = OxmlElement("w:" + side)
            borders.append(node)
        node.set(tag("val"), "single")
        node.set(tag("sz"), "4")
        node.set(tag("space"), "0")
        node.set(tag("color"), "D9D9D9")


def set_table_margins(table) -> None:
    tbl_pr = table._tbl.tblPr
    margins = find(tbl_pr, "tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side, value in (("top", 80), ("left", 100), ("bottom", 80), ("right", 100)):
        node = find(margins, side)
        if node is None:
            node = OxmlElement("w:" + side)
            margins.append(node)
        node.set(tag("w"), str(value))
        node.set(tag("type"), "dxa")


def set_table_widths(table, available_twips: int) -> None:
    columns = len(table.columns)
    if columns == 0:
        return
    scores = [table_column_score(table, index) for index in range(columns)]
    minimum = max(720, min(1050, available_twips // (columns * 4)))
    remaining = max(0, available_twips - minimum * columns)
    total_score = sum(scores) or columns
    widths = [minimum + round(remaining * score / total_score) for score in scores]
    widths[-1] += available_twips - sum(widths)

    tbl_pr = table._tbl.tblPr
    tbl_w = find(tbl_pr, "tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(tag("w"), str(available_twips))
    tbl_w.set(tag("type"), "dxa")
    layout = find(tbl_pr, "tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(tag("type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(tag("w"), str(width))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            if index < len(widths):
                set_cell_width(cell, widths[index])


def mark_header_row(table) -> None:
    if len(table.rows) < 2:
        return
    row = table.rows[0]._tr
    tr_pr = row.get_or_add_trPr()
    remove_all(tr_pr, "tblHeader")
    tr_pr.append(OxmlElement("w:tblHeader"))


def remove_fixed_row_heights(table) -> None:
    for row in table.rows:
        tr_pr = row._tr.find(tag("trPr"))
        if tr_pr is not None:
            remove_all(tr_pr, "trHeight")
            remove_all(tr_pr, "cantSplit")


def all_tables(document: Document) -> list:
    result = []

    def collect(tables) -> None:
        for table in tables:
            result.append(table)
            for row in table.rows:
                for cell in row.cells:
                    collect(cell.tables)

    collect(document.tables)
    return result


def format_tables(document: Document) -> None:
    if not document.sections:
        return
    section = document.sections[0]
    available = int(round((section.page_width - section.left_margin - section.right_margin) / 635))
    for table in all_tables(document):
        tbl_pr = table._tbl.tblPr
        for child_name in ("tblStyle", "tblLook", "tblInd", "tblCellSpacing", "tblpPr"):
            remove_all(tbl_pr, child_name)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_widths(table, available)
        set_table_borders(table)
        set_table_margins(table)
        mark_header_row(table)
        remove_fixed_row_heights(table)
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    role = "Code Block" if is_technical(clean(paragraph.text), paragraph.text) else "Table Text"
                    style = document.styles[role]
                    clean_paragraph_properties(paragraph, role, style.style_id)
                    clean_run_properties(paragraph)


def scale_inline_images(document: Document) -> None:
    for section in document.sections:
        available = section.page_width - section.left_margin - section.right_margin
        for shape in document.inline_shapes:
            if shape.width <= available:
                continue
            ratio = available / shape.width
            shape.width = int(shape.width * ratio)
            shape.height = int(shape.height * ratio)


def format_document(input_path: Path, output_path: Path, profile: str) -> None:
    if input_path.resolve() == output_path.resolve():
        raise ValueError("输出文件必须与输入文件不同，原文档不会被覆盖")
    source = Document(str(input_path))
    resolved_profile = detect_profile(source) if profile == "auto" else profile
    original_paragraphs = meaningful_paragraphs(source)
    roles = classify_paragraphs(original_paragraphs, resolved_profile)

    if not BASE_TEMPLATE_PATH.is_file():
        raise FileNotFoundError(f"找不到空白基线模板：{BASE_TEMPLATE_PATH}")
    document = Document(str(BASE_TEMPLATE_PATH))
    clear_main_body(document)
    configure_styles(document, reset_numbering=True)
    configure_document(document, resolved_profile)

    content = source_body_content(source)
    target_body = document._body._element
    for node in content:
        target_body.insert(len(target_body) - 1, node)
    relationship_map = copy_source_relationships(source, document, content)
    for node in content:
        remap_relationship_ids(node, relationship_map)

    target_paragraphs = all_paragraphs(document)
    if len(target_paragraphs) != len(roles):
        raise ValueError(
            f"迁移后段落数量不一致：源文档 {len(roles)}，目标文档 {len(target_paragraphs)}"
        )

    for paragraph, (_, role) in zip(target_paragraphs, roles):
        if role not in {style.name for style in document.styles}:
            role = "Normal"
        strip_heading_prefix(paragraph, role)
        if role != "Code Block" and role not in {"Equation", "TOC 1", "TOC 2", "TOC 3", "Figure List", "Table List"}:
            strip_paragraph_edge_whitespace(paragraph)
        style = document.styles[role]
        clean_paragraph_properties(paragraph, role, style.style_id)
        clean_run_properties(paragraph)

    if resolved_profile == "thesis":
        copy_thesis_headers(source, document)
    format_tables(document)
    scale_inline_images(document)
    add_baseline_marker(document)
    document.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="将 DOCX 迁移到标准中文文档格式并输出新文件")
    parser.add_argument("input", type=Path, help="输入 DOCX，只读")
    parser.add_argument("--output", required=True, type=Path, help="输出 DOCX，不得与输入相同")
    parser.add_argument("--profile", choices=("auto", "thesis", "general"), default="auto")
    args = parser.parse_args()
    if args.input.suffix.lower() != ".docx" or not args.input.is_file():
        raise SystemExit("输入必须是存在的 .docx 文件")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    format_document(args.input.resolve(), args.output.resolve(), args.profile)
    print(f"已生成：{args.output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
