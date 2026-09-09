#!/usr/bin/env python3
"""生成干净的中文 DOCX 基线模板。

模板只包含页面、节、样式、页码和编号定义，不包含任何业务正文。
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def tag(name: str) -> str:
    return qn("w:" + name)


def find(parent, name: str):
    return parent.find(tag(name)) if parent is not None else None


def ensure(parent, name: str):
    node = find(parent, name)
    if node is None:
        node = OxmlElement("w:" + name)
        parent.append(node)
    return node


def remove_all(parent, name: str) -> None:
    if parent is None:
        return
    for node in list(parent.findall(tag(name))):
        parent.remove(node)


def remove_attr(node, name: str) -> None:
    if node is not None:
        node.attrib.pop(tag(name), None)


def get_style_by_id(document: Document, style_id: str):
    for style in document.styles:
        if style.style_id == style_id:
            return style
    return None


def get_or_add_style(document: Document, name: str, style_type):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, style_type)


def reset_style_links(style, *, based_on: str | None = None, next_style: str | None = None, link: str | None = None) -> None:
    element = style._element
    for child_name in ("basedOn", "next", "link"):
        remove_all(element, child_name)
    for child_name, value in (("basedOn", based_on), ("next", next_style), ("link", link)):
        if value:
            node = OxmlElement("w:" + child_name)
            node.set(tag("val"), value)
            element.insert(0, node)


def set_style_fonts(
    style,
    east_asia: str,
    latin: str,
    size_pt: float,
    *,
    bold: bool = False,
    italic: bool = False,
    no_proof: bool = False,
) -> None:
    """显式写入四类字体，避免主题字体和旧字符属性回流。"""
    style.font.name = latin
    style.font.size = Pt(size_pt)
    rpr = style._element.get_or_add_rPr()

    for child_name in (
        "rStyle",
        "b",
        "i",
        "caps",
        "smallCaps",
        "strike",
        "outline",
        "shadow",
        "emboss",
        "imprint",
        "vanish",
        "u",
        "color",
        "sz",
        "szCs",
        "rFonts",
        "noProof",
    ):
        remove_all(rpr, child_name)

    rfonts = OxmlElement("w:rFonts")
    for name, value in {
        "ascii": latin,
        "hAnsi": latin,
        "eastAsia": east_asia,
        "cs": latin,
    }.items():
        rfonts.set(tag(name), value)
    rpr.insert(0, rfonts)

    size = OxmlElement("w:sz")
    size.set(tag("val"), str(round(size_pt * 2)))
    rpr.append(size)

    color = OxmlElement("w:color")
    color.set(tag("val"), "000000")
    rpr.append(color)

    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    if no_proof:
        rpr.append(OxmlElement("w:noProof"))


def set_style_paragraph(
    style,
    *,
    alignment: int,
    line_twips: int,
    before_twips: int | None = 0,
    after_twips: int | None = 0,
    after_lines: int | None = None,
    first_line_chars: int | None = None,
    first_line_twips: int | None = None,
    hanging_twips: int | None = None,
    keep_next: bool = False,
    keep_lines: bool = False,
    page_break_before: bool = False,
    outline_level: int | None = None,
) -> None:
    ppr = style._element.get_or_add_pPr()
    for child_name in (
        "jc",
        "spacing",
        "ind",
        "keepNext",
        "keepLines",
        "pageBreakBefore",
        "widowControl",
        "outlineLvl",
        "tabs",
        "pBdr",
        "shd",
        "numPr",
    ):
        remove_all(ppr, child_name)

    jc = OxmlElement("w:jc")
    jc.set(tag("val"), {0: "left", 1: "center", 2: "right", 3: "both"}[alignment])
    ppr.append(jc)

    spacing = OxmlElement("w:spacing")
    if before_twips is not None:
        spacing.set(tag("before"), str(before_twips))
    if after_twips is not None:
        spacing.set(tag("after"), str(after_twips))
    if after_lines is not None:
        spacing.set(tag("afterLines"), str(after_lines))
    spacing.set(tag("line"), str(line_twips))
    spacing.set(tag("lineRule"), "auto")
    ppr.append(spacing)

    indent = OxmlElement("w:ind")
    indent.set(tag("left"), "0")
    indent.set(tag("right"), "0")
    if first_line_chars is not None:
        indent.set(tag("firstLineChars"), str(first_line_chars))
    if first_line_twips is not None:
        indent.set(tag("firstLine"), str(first_line_twips))
    if hanging_twips is not None:
        indent.set(tag("hanging"), str(hanging_twips))
    elif first_line_chars is None and first_line_twips is None:
        # 非正文样式显式清零首行缩进，避免从 Normal 继承两个字符。
        indent.set(tag("firstLine"), "0")
    ppr.append(indent)

    if keep_next:
        ppr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        ppr.append(OxmlElement("w:keepLines"))
    if page_break_before:
        ppr.append(OxmlElement("w:pageBreakBefore"))
    if outline_level is not None:
        outline = OxmlElement("w:outlineLvl")
        outline.set(tag("val"), str(outline_level))
        ppr.append(outline)


def set_style_meta(style, *, based_on: str | None = None, next_style: str | None = None, link: str | None = None) -> None:
    reset_style_links(style, based_on=based_on, next_style=next_style, link=link)
    for name in ("pBdr", "shd"):
        remove_all(style._element.get_or_add_pPr(), name)


def ensure_character_style(document: Document, style_id: str, name: str, east_asia: str, latin: str, size: float, bold: bool = False) -> str:
    character = get_style_by_id(document, style_id)
    if character is None:
        character = get_or_add_style(document, name, WD_STYLE_TYPE.CHARACTER)
        character._element.set(tag("styleId"), style_id)
    set_style_meta(character)
    set_style_fonts(character, east_asia, latin, size, bold=bold)
    return character.style_id


def set_linked_character_style(document: Document, paragraph_style, name: str, east_asia: str, latin: str, size: float, bold: bool) -> None:
    style_id = paragraph_style.style_id.replace(" ", "") + "Char"
    character_id = ensure_character_style(document, style_id, name, east_asia, latin, size, bold)
    normal_id = document.styles["Normal"].style_id
    reset_style_links(paragraph_style, based_on=normal_id, next_style=normal_id, link=character_id)


def configure_doc_defaults(document: Document) -> None:
    styles_root = document.styles._element
    doc_defaults = ensure(styles_root, "docDefaults")
    rpr_default = ensure(doc_defaults, "rPrDefault")
    rpr = ensure(rpr_default, "rPr")
    for child_name in ("rFonts", "sz", "szCs", "color", "b", "i", "u"):
        remove_all(rpr, child_name)
    rfonts = OxmlElement("w:rFonts")
    for name, value in {"ascii": "Times New Roman", "hAnsi": "Times New Roman", "eastAsia": "宋体", "cs": "Times New Roman"}.items():
        rfonts.set(tag(name), value)
    rpr.append(rfonts)
    size = OxmlElement("w:sz")
    size.set(tag("val"), "24")
    rpr.append(size)
    color = OxmlElement("w:color")
    color.set(tag("val"), "000000")
    rpr.append(color)

    ppr_default = ensure(doc_defaults, "pPrDefault")
    ppr = ensure(ppr_default, "pPr")
    for child_name in ("spacing", "ind", "jc", "keepNext", "keepLines", "pageBreakBefore"):
        remove_all(ppr, child_name)


def add_numbering(document: Document, *, reset: bool = False) -> tuple[int, int]:
    """创建主文档和附录编号；isLgl 保证二级编号使用阿拉伯章号。"""
    numbering = document.part.numbering_part.element
    if reset:
        for child_name in ("abstractNum", "num"):
            remove_all(numbering, child_name)

    abstract_ids = [
        int(node.get(tag("abstractNumId")))
        for node in numbering.findall(tag("abstractNum"))
        if node.get(tag("abstractNumId"))
    ]
    num_ids = [
        int(node.get(tag("numId")))
        for node in numbering.findall(tag("num"))
        if node.get(tag("numId"))
    ]
    main_abstract_id = 1 if reset else max(abstract_ids, default=0) + 1
    appendix_abstract_id = main_abstract_id + 1
    main_num_id = 1 if reset else max(num_ids, default=0) + 1
    appendix_num_id = main_num_id + 1

    def add_level(abstract, level: int, number_format: str, text: str, style_id: str, alignment: str, legal: bool = False) -> None:
        item = OxmlElement("w:lvl")
        item.set(tag("ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(tag("val"), "1")
        item.append(start)
        if legal:
            item.append(OxmlElement("w:isLgl"))
        fmt = OxmlElement("w:numFmt")
        fmt.set(tag("val"), number_format)
        item.append(fmt)
        label = OxmlElement("w:lvlText")
        label.set(tag("val"), text)
        item.append(label)
        suffix = OxmlElement("w:suff")
        suffix.set(tag("val"), "space")
        item.append(suffix)
        justify = OxmlElement("w:lvlJc")
        justify.set(tag("val"), alignment)
        item.append(justify)
        ppr = OxmlElement("w:pPr")
        list_indent = OxmlElement("w:ind")
        list_indent.set(tag("left"), "0")
        list_indent.set(tag("hanging"), "0")
        list_indent.set(tag("firstLine"), "0")
        ppr.append(list_indent)
        pstyle = OxmlElement("w:pStyle")
        pstyle.set(tag("val"), style_id)
        ppr.append(pstyle)
        item.append(ppr)
        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(tag("ascii"), "Times New Roman")
        rfonts.set(tag("hAnsi"), "Times New Roman")
        rfonts.set(tag("eastAsia"), "黑体")
        rpr.append(rfonts)
        color = OxmlElement("w:color")
        color.set(tag("val"), "000000")
        rpr.append(color)
        item.append(rpr)
        abstract.append(item)

    main = OxmlElement("w:abstractNum")
    main.set(tag("abstractNumId"), str(main_abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(tag("val"), "multilevel")
    main.append(multi)
    add_level(main, 0, "chineseCounting", "第%1章", "Heading1", "center")
    add_level(main, 1, "decimal", "%1.%2", "Heading2", "left", legal=True)
    add_level(main, 2, "decimal", "%1.%2.%3", "Heading3", "left", legal=True)

    appendix = OxmlElement("w:abstractNum")
    appendix.set(tag("abstractNumId"), str(appendix_abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(tag("val"), "singleLevel")
    appendix.append(multi)
    add_level(appendix, 0, "decimal", "附录%1", "AppendixHeading", "center")
    numbering.extend([main, appendix])

    def add_num(num_id: int, abstract_id: int) -> None:
        num = OxmlElement("w:num")
        num.set(tag("numId"), str(num_id))
        reference = OxmlElement("w:abstractNumId")
        reference.set(tag("val"), str(abstract_id))
        num.append(reference)
        numbering.append(num)

    add_num(main_num_id, main_abstract_id)
    add_num(appendix_num_id, appendix_abstract_id)
    return main_num_id, appendix_num_id


def bind_numbering(style, level: int, num_id: int) -> None:
    ppr = style._element.get_or_add_pPr()
    remove_all(ppr, "numPr")
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(tag("val"), str(level))
    num = OxmlElement("w:numId")
    num.set(tag("val"), str(num_id))
    num_pr.extend([ilvl, num])
    ppr.append(num_pr)


def configure_styles(document: Document, *, reset_numbering: bool = False) -> None:
    configure_doc_defaults(document)

    normal = document.styles["Normal"]
    normal_id = normal.style_id
    set_style_meta(normal)
    set_style_fonts(normal, "宋体", "Times New Roman", 12)
    set_style_paragraph(normal, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, line_twips=360, first_line_chars=200, first_line_twips=480)

    footer = document.styles["Footer"]
    set_style_meta(footer, based_on=normal_id, next_style=normal_id)
    set_style_fonts(footer, "宋体", "Times New Roman", 10.5)
    set_style_paragraph(footer, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_twips=240)

    header = document.styles["Header"]
    set_style_meta(header, based_on=normal_id, next_style=normal_id)
    set_style_fonts(header, "宋体", "Times New Roman", 10.5)
    set_style_paragraph(header, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_twips=240)

    title = document.styles["Title"]
    set_style_meta(title, based_on=normal_id, next_style=normal_id)
    set_style_fonts(title, "黑体", "Times New Roman", 16, bold=False)
    set_style_paragraph(title, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_twips=360, keep_next=True, keep_lines=True)
    set_linked_character_style(document, title, "Title Char", "黑体", "Times New Roman", 16, False)

    heading_specs = (
        ("Heading 1", 16, WD_ALIGN_PARAGRAPH.CENTER, 0, 100, True),
        ("Heading 2", 14, WD_ALIGN_PARAGRAPH.LEFT, 1, None, False),
        ("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT, 2, None, False),
    )
    for name, size, alignment, outline, after_lines, page_break in heading_specs:
        style = document.styles[name]
        set_style_meta(style, based_on=normal_id, next_style=normal_id)
        set_style_fonts(style, "黑体", "Times New Roman", size, bold=False)
        set_style_paragraph(
            style,
            alignment=alignment,
            line_twips=360,
            after_twips=None if after_lines is not None else 0,
            after_lines=after_lines,
            keep_next=True,
            keep_lines=True,
            page_break_before=page_break,
            outline_level=outline,
        )
        set_linked_character_style(document, style, name + " Char", "黑体", "Times New Roman", size, False)

    figure = get_or_add_style(document, "Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    set_style_meta(figure, based_on=normal_id, next_style=normal_id)
    set_style_fonts(figure, "楷体", "Times New Roman", 10.5)
    set_style_paragraph(figure, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_twips=300, keep_next=True, keep_lines=True)

    table_caption = get_or_add_style(document, "Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    set_style_meta(table_caption, based_on=normal_id, next_style=normal_id)
    set_style_fonts(table_caption, "楷体", "Times New Roman", 10.5)
    set_style_paragraph(table_caption, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_twips=300, keep_next=True, keep_lines=True)

    equation = get_or_add_style(document, "Equation", WD_STYLE_TYPE.PARAGRAPH)
    set_style_meta(equation, based_on=normal_id, next_style=normal_id)
    set_style_fonts(equation, "Cambria Math", "Cambria Math", 12)
    set_style_paragraph(equation, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_twips=240, keep_next=True, keep_lines=True)
    equation.paragraph_format.tab_stops.clear_all()
    equation.paragraph_format.tab_stops.add_tab_stop(Cm(8.25), WD_TAB_ALIGNMENT.CENTER, WD_TAB_LEADER.SPACES)
    equation.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)

    reference = get_or_add_style(document, "Reference", WD_STYLE_TYPE.PARAGRAPH)
    set_style_meta(reference, based_on=normal_id, next_style=normal_id)
    set_style_fonts(reference, "宋体", "Times New Roman", 12)
    set_style_paragraph(reference, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_twips=360, hanging_twips=480)

    abstract_body = get_or_add_style(document, "Abstract Body", WD_STYLE_TYPE.PARAGRAPH)
    set_style_meta(abstract_body, based_on=normal_id, next_style=normal_id)
    set_style_fonts(abstract_body, "宋体", "Times New Roman", 14)
    set_style_paragraph(abstract_body, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, line_twips=360, first_line_chars=200, first_line_twips=560)

    english_body = get_or_add_style(document, "English Abstract Body", WD_STYLE_TYPE.PARAGRAPH)
    set_style_meta(english_body, based_on=normal_id, next_style=normal_id)
    set_style_fonts(english_body, "Times New Roman", "Times New Roman", 14)
    set_style_paragraph(english_body, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, line_twips=360, first_line_chars=400, first_line_twips=560)

    for name in ("Abstract Heading", "English Abstract Heading", "TOC Title"):
        style = get_or_add_style(document, name, WD_STYLE_TYPE.PARAGRAPH)
        set_style_meta(style, based_on=normal_id, next_style=normal_id)
        set_style_fonts(style, "黑体", "Times New Roman", 16, bold=False)
        set_style_paragraph(style, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_twips=360, keep_next=True, keep_lines=True)

    for level in range(1, 4):
        style = get_or_add_style(document, f"TOC {level}", WD_STYLE_TYPE.PARAGRAPH)
        set_style_meta(style, based_on=normal_id, next_style=normal_id)
        set_style_fonts(style, "宋体", "Times New Roman", 12)
        set_style_paragraph(style, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_twips=300)

    for name in ("Figure List", "Table List"):
        style = get_or_add_style(document, name, WD_STYLE_TYPE.PARAGRAPH)
        set_style_meta(style, based_on=normal_id, next_style=normal_id)
        set_style_fonts(style, "楷体", "Times New Roman", 10.5)
        set_style_paragraph(style, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_twips=300)

    appendix = get_or_add_style(document, "Appendix Heading", WD_STYLE_TYPE.PARAGRAPH)
    set_style_meta(appendix, based_on=normal_id, next_style=normal_id)
    set_style_fonts(appendix, "黑体", "Times New Roman", 16, bold=False)
    set_style_paragraph(appendix, alignment=WD_ALIGN_PARAGRAPH.CENTER, line_twips=360, keep_next=True, keep_lines=True, page_break_before=True)

    code = get_or_add_style(document, "Code Block", WD_STYLE_TYPE.PARAGRAPH)
    set_style_meta(code, based_on=normal_id, next_style=normal_id)
    set_style_fonts(code, "宋体", "Consolas", 10.5, no_proof=True)
    set_style_paragraph(code, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_twips=240, after_twips=60, keep_lines=True)

    table_text = get_or_add_style(document, "Table Text", WD_STYLE_TYPE.PARAGRAPH)
    set_style_meta(table_text, based_on=normal_id, next_style=normal_id)
    set_style_fonts(table_text, "宋体", "Times New Roman", 10.5)
    set_style_paragraph(table_text, alignment=WD_ALIGN_PARAGRAPH.LEFT, line_twips=300)

    for name, east_asia, size, bold in (
        ("Keyword Label", "黑体", 14, False),
        ("Keyword Content", "宋体", 14, False),
    ):
        character = get_or_add_style(document, name, WD_STYLE_TYPE.CHARACTER)
        set_style_meta(character)
        set_style_fonts(character, east_asia, "Times New Roman", size, bold=bold)

    main_num_id, appendix_num_id = add_numbering(document, reset=reset_numbering)
    bind_numbering(document.styles["Heading 1"], 0, main_num_id)
    bind_numbering(document.styles["Heading 2"], 1, main_num_id)
    bind_numbering(document.styles["Heading 3"], 2, main_num_id)
    bind_numbering(appendix, 0, appendix_num_id)


def clear_container(container, style_id: str = "Normal") -> None:
    paragraphs = list(container.paragraphs)
    paragraph = paragraphs[0]
    for extra in paragraphs[1:]:
        container._element.remove(extra._p)
    ppr = paragraph._p.get_or_add_pPr()
    for node in list(ppr):
        if node.tag != tag("sectPr"):
            ppr.remove(node)
    for node in list(paragraph._p):
        if node.tag != tag("pPr"):
            paragraph._p.remove(node)
    pstyle = OxmlElement("w:pStyle")
    pstyle.set(tag("val"), style_id)
    ppr.insert(0, pstyle)


def add_page_field(section, footer_style_id: str) -> None:
    footer = section.footer
    clear_container(footer, footer_style_id)
    paragraph = footer.paragraphs[0]
    field = OxmlElement("w:fldSimple")
    field.set(tag("instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for name, value in {"ascii": "Times New Roman", "hAnsi": "Times New Roman", "eastAsia": "Times New Roman", "cs": "Times New Roman"}.items():
        rfonts.set(tag(name), value)
    rpr.append(rfonts)
    size = OxmlElement("w:sz")
    size.set(tag("val"), "21")
    rpr.append(size)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def add_baseline_marker(document: Document) -> None:
    settings = document.settings._element
    update = find(settings, "updateFields")
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(tag("val"), "true")
    variables = ensure(settings, "docVars")
    for item in list(variables.findall(tag("docVar"))):
        if item.get(tag("name")) == "chinese-docx-format-base":
            variables.remove(item)
    marker = OxmlElement("w:docVar")
    marker.set(tag("name"), "chinese-docx-format-base")
    marker.set(tag("val"), "1")
    variables.append(marker)


def set_setting_flag(document: Document, name: str, enabled: bool) -> None:
    """设置文档级开关，避免继承源文档的奇偶页眉配置。"""
    settings = document.settings._element
    remove_all(settings, name)
    if enabled:
        settings.append(OxmlElement("w:" + name))


def configure_document(document: Document, profile: str = "general") -> None:
    if profile not in {"general", "thesis"}:
        raise ValueError(f"不支持的文档模式：{profile}")

    thesis = profile == "thesis"
    normal_id = document.styles["Normal"].style_id
    header_id = document.styles["Header"].style_id
    footer_id = document.styles["Footer"].style_id
    set_setting_flag(document, "evenAndOddHeaders", thesis)
    remove_all(document.settings._element, "characterSpacingControl")
    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        section.bottom_margin = Cm(2.7)
        section.header_distance = Cm(1.5 if thesis else 0.8)
        section.footer_distance = Cm(1.75 if thesis else 1.0)
        sect_pr = section._sectPr
        remove_all(sect_pr, "titlePg")
        if thesis:
            sect_pr.append(OxmlElement("w:titlePg"))
        page_number = find(sect_pr, "pgNumType")
        if page_number is None:
            page_number = OxmlElement("w:pgNumType")
            sect_pr.append(page_number)
        page_number.set(tag("start"), "1")
        remove_attr(page_number, "fmt")
        add_page_field(section, footer_id)
        clear_container(section.header, header_id)
        clear_container(section.first_page_header, header_id)
        clear_container(section.even_page_header, header_id)

    add_baseline_marker(document)
    properties = document.core_properties
    for name in ("author", "last_modified_by", "title", "subject", "keywords", "comments"):
        setattr(properties, name, "")


def build(output: Path, profile: str = "general") -> None:
    document = Document()
    configure_styles(document, reset_numbering=True)
    configure_document(document, profile)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="生成干净的中文 DOCX 基线模板")
    parser.add_argument("output", type=Path, help="输出 DOCX 路径")
    parser.add_argument("--profile", choices=("general", "thesis"), default="general")
    args = parser.parse_args()
    build(args.output.resolve(), args.profile)
    print(f"已生成：{args.output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
