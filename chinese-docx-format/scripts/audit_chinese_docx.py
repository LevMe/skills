#!/usr/bin/env python3
"""只读审计中文 DOCX 的页面、样式、结构和常见格式问题。"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from zipfile import BadZipFile, ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W = f"{{{W_NS}}}"
WP = f"{{{WP_NS}}}"
M = f"{{{M_NS}}}"

CM_PER_INCH = 2.54
EMU_PER_INCH = 914400
TWIPS_PER_INCH = 1440
TOLERANCE_CM = 0.04
BASELINE_MARKER_NAME = "chinese-docx-format-base"
BASELINE_MARKER_VALUE = "1"


def clean(text: str | None) -> str:
    """压缩可见文本中的空白，便于结构匹配。"""
    return re.sub(r"\s+", " ", text or "").strip()


def add_issue(report: dict, severity: str, code: str, message: str, details: str | None = None) -> None:
    item = {"code": code, "message": message}
    if details:
        item["details"] = details
    report[severity].append(item)


def record_check(report: dict, code: str, passed: bool, message: str) -> None:
    report["checks"].append({"code": code, "passed": passed, "message": message})


def child(parent, tag: str):
    return parent.find(W + tag) if parent is not None else None


def attribute(element, name: str) -> str | None:
    return element.get(W + name) if element is not None else None


def paragraph_style_role(paragraph) -> str | None:
    style_id = paragraph.style.style_id if paragraph.style else ""
    style_name = paragraph.style.name if paragraph.style else ""
    aliases = {
        "Heading1": "Heading 1",
        "Heading2": "Heading 2",
        "Heading3": "Heading 3",
        "标题 1": "Heading 1",
        "标题 2": "Heading 2",
        "标题 3": "Heading 3",
        "Figure Caption": "Figure Caption",
        "Table Caption": "Table Caption",
        "Figure List": "Figure List",
        "Table List": "Table List",
        "图题": "Figure Caption",
        "表题": "Table Caption",
        "Reference": "Reference",
        "EndNote Bibliography": "Reference",
        "Abstract Body": "Abstract Body",
        "English Abstract Body": "English Abstract Body",
        "Appendix Heading": "Appendix Heading",
    }
    if style_id in aliases:
        return aliases[style_id]
    return aliases.get(style_name, style_name)


def find_style(doc, names: set[str], ids: set[str] | None = None):
    ids = ids or set()
    for style in doc.styles:
        if style.name in names or style.style_id in ids:
            return style
    return None


def style_xml(style):
    return style._element


def style_rpr(style):
    return style_xml(style).find(W + "rPr")


def style_ppr(style):
    return style_xml(style).find(W + "pPr")


def style_fonts(style) -> dict[str, str | None]:
    rpr = style_rpr(style)
    rfonts = rpr.find(W + "rFonts") if rpr is not None else None
    if rfonts is None:
        return {"eastAsia": None, "ascii": None, "hAnsi": None}
    return {
        "eastAsia": rfonts.get(W + "eastAsia"),
        "ascii": rfonts.get(W + "ascii"),
        "hAnsi": rfonts.get(W + "hAnsi"),
    }


def style_chain(doc, style) -> tuple[list, bool, str | None]:
    """解析 basedOn 继承链，并返回循环或悬空继承信息。"""
    chain = []
    seen: set[str] = set()
    current = style
    while current is not None:
        style_id = current.style_id
        if style_id in seen:
            return chain, True, None
        seen.add(style_id)
        chain.append(current)
        based_on = child(style_xml(current), "basedOn")
        base_id = attribute(based_on, "val")
        if not base_id:
            return chain, False, None
        current = find_style(doc, set(), {base_id})
        if current is None:
            return chain, False, base_id
    return chain, False, None


def effective_style_fonts(doc, style) -> dict[str, str | None]:
    """按最近样式优先解析最终字体，避免漏检 basedOn 中的字体。"""
    result = {"eastAsia": None, "ascii": None, "hAnsi": None}
    chain, _, _ = style_chain(doc, style)
    for current in chain:
        for key, value in style_fonts(current).items():
            if result[key] is None and value:
                result[key] = value
    return result


def effective_style_decoration_issues(doc, style) -> list[str]:
    """按继承链解析标题最终颜色、下划线、边框、底纹和主题字体。"""
    chain, _, _ = style_chain(doc, style)
    color = None
    underline = None
    borders = None
    shading = None
    theme_fonts: dict[str, str] = {}
    for current in chain:
        rpr = style_rpr(current)
        if rpr is not None:
            if color is None:
                color = attribute(child(rpr, "color"), "val")
            if underline is None:
                underline = attribute(child(rpr, "u"), "val")
            fonts = rpr.find(W + "rFonts")
            if fonts is not None:
                for name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                    if name not in theme_fonts and attribute(fonts, name):
                        theme_fonts[name] = attribute(fonts, name)
        ppr = style_ppr(current)
        if ppr is not None:
            if borders is None:
                borders = child(ppr, "pBdr")
            if shading is None:
                shading = child(ppr, "shd")

    issues: list[str] = []
    if color not in {None, "000000", "auto"}:
        issues.append(f"颜色={color}")
    if underline not in {None, "none"}:
        issues.append(f"下划线={underline}")
    if borders is not None:
        visible = [attribute(item, "val") for item in list(borders) if attribute(item, "val") not in {None, "none"}]
        if visible:
            issues.append("存在段落边框")
    if shading is not None and attribute(shading, "fill") not in {None, "auto", "FFFFFF"}:
        issues.append(f"底纹={attribute(shading, 'fill')}")
    if theme_fonts:
        issues.append("存在主题字体属性")
    return issues


def style_size_pt(style) -> float | None:
    rpr = style_rpr(style)
    size = rpr.find(W + "sz") if rpr is not None else None
    value = attribute(size, "val")
    return int(value) / 2 if value and value.isdigit() else None


def style_alignment(style) -> int | None:
    ppr = style_ppr(style)
    alignment = ppr.find(W + "jc") if ppr is not None else None
    value = attribute(alignment, "val")
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }.get(value)


def style_line_twips(style) -> str | None:
    ppr = style_ppr(style)
    spacing = ppr.find(W + "spacing") if ppr is not None else None
    return attribute(spacing, "line")


def style_spacing(style) -> dict[str, str | None]:
    ppr = style_ppr(style)
    spacing = ppr.find(W + "spacing") if ppr is not None else None
    return {
        name: attribute(spacing, name)
        for name in ("before", "after", "afterLines", "line", "lineRule")
    }


def style_indent(style) -> tuple[str | None, str | None]:
    ppr = style_ppr(style)
    indent = ppr.find(W + "ind") if ppr is not None else None
    return attribute(indent, "firstLineChars"), attribute(indent, "firstLine")


def effective_style_indent(doc, style) -> dict[str, str | None]:
    """按继承链解析段落缩进，避免标题继承 Normal 的首行缩进。"""
    result = {name: None for name in ("left", "right", "firstLine", "firstLineChars", "hanging")}
    first_line_seen = False
    chain, _, _ = style_chain(doc, style)
    for current in chain:
        ppr = style_ppr(current)
        indent = ppr.find(W + "ind") if ppr is not None else None
        if indent is None:
            continue
        for name in ("left", "right"):
            if result[name] is None:
                result[name] = attribute(indent, name)
        if first_line_seen:
            continue
        for name in ("firstLine", "firstLineChars", "hanging"):
            value = attribute(indent, name)
            if value is not None:
                result[name] = value
                first_line_seen = True
                break
    return result


def effective_style_has_tabs(doc, style) -> bool:
    """检查标题样式继承链中是否残留制表位定义。"""
    chain, _, _ = style_chain(doc, style)
    return any(
        style_ppr(current) is not None and style_ppr(current).find(W + "tabs") is not None
        for current in chain
    )


def style_boolean(doc, style, name: str) -> bool:
    """按继承链解析开关属性，显式 false 也必须覆盖父样式。"""
    chain, _, _ = style_chain(doc, style)
    for current in chain:
        node = child(style_ppr(current), name)
        if node is None:
            continue
        return attribute(node, "val") not in {"0", "false", "off", "no"}
    return False


def style_run_boolean(doc, style, name: str) -> bool:
    """解析字符属性中的加粗、倾斜和不检查拼写等开关。"""
    chain, _, _ = style_chain(doc, style)
    for current in chain:
        node = child(style_rpr(current), name)
        if node is None:
            continue
        return attribute(node, "val") not in {"0", "false", "off", "no"}
    return False


def style_linked_character(doc, style):
    """返回段落样式关联的字符样式。"""
    link = child(style_xml(style), "link")
    link_id = attribute(link, "val")
    return find_style(doc, set(), {link_id}) if link_id else None


def style_numpr(style) -> tuple[str | None, str | None]:
    ppr = style_ppr(style)
    numpr = ppr.find(W + "numPr") if ppr is not None else None
    if numpr is None:
        return None, None
    ilvl = numpr.find(W + "ilvl")
    num_id = numpr.find(W + "numId")
    return attribute(ilvl, "val"), attribute(num_id, "val")


def paragraph_in_table(paragraph) -> bool:
    """判断段落是否位于表格单元格中。"""
    node = paragraph._p
    while node is not None:
        if node.tag == W + "tc":
            return True
        node = node.getparent()
    return False


def paragraph_has_field(paragraph, token: str | None = None) -> bool:
    instructions = " ".join((node.text or "") for node in paragraph._p.iter(W + "instrText"))
    instructions += " " + " ".join(attribute(node, "instr") or "" for node in paragraph._p.iter(W + "fldSimple"))
    if token is None:
        return bool(instructions.strip())
    return token.upper() in instructions.upper()


def append_table_paragraphs(result: list, tables) -> None:
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                append_table_paragraphs(result, cell.tables)
                for paragraph in cell.paragraphs:
                    if paragraph not in result:
                        result.append(paragraph)


def all_paragraphs(doc) -> list:
    """收集正文、表格、页眉和页脚中的段落。"""
    result = list(doc.paragraphs)
    append_table_paragraphs(result, doc.tables)
    containers = []
    for section in doc.sections:
        containers.extend(
            [
                section.header,
                section.first_page_header,
                section.even_page_header,
                section.footer,
                section.first_page_footer,
                section.even_page_footer,
            ]
        )
    seen_parts: set[int] = set()
    for container in containers:
        part_key = id(container._element)
        if part_key in seen_parts:
            continue
        seen_parts.add(part_key)
        result.extend(container.paragraphs)
        append_table_paragraphs(result, container.tables)
    return result


def all_tables(doc) -> list:
    result: list = []

    def collect(tables) -> None:
        for table in tables:
            result.append(table)
            for row in table.rows:
                for cell in row.cells:
                    collect(cell.tables)

    collect(doc.tables)
    for section in doc.sections:
        for container in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            collect(container.tables)
    return result


def check_style(
    report: dict,
    doc,
    name: str,
    ids: set[str],
    *,
    east_asia: str | None = None,
    latin: str | None = None,
    size: float | None = None,
    alignment: int | None = None,
    line: str | None = None,
    first_line_chars: str | None = None,
    require_num_level: str | None = None,
) -> object | None:
    style = find_style(doc, {name}, ids)
    if style is None:
        add_issue(report, "errors", "STYLE_MISSING", f"缺少必需样式：{name}")
        return None

    fonts = effective_style_fonts(doc, style)
    if east_asia and fonts["eastAsia"] != east_asia:
        add_issue(report, "errors", "STYLE_FONT_EAST_ASIA", f"{name} 的中文字体不是 {east_asia}", str(fonts["eastAsia"]))
    if latin and fonts["ascii"] != latin and fonts["hAnsi"] != latin:
        add_issue(report, "errors", "STYLE_FONT_LATIN", f"{name} 的西文字体不是 {latin}", str(fonts))
    if size is not None:
        actual = style_size_pt(style)
        if actual is None or abs(actual - size) > 0.01:
            add_issue(report, "errors", "STYLE_SIZE", f"{name} 的字号不是 {size:g} pt", str(actual))
    if alignment is not None and style_alignment(style) != alignment:
        add_issue(report, "errors", "STYLE_ALIGNMENT", f"{name} 的对齐方式不符合规范")
    if line is not None and style_line_twips(style) != line:
        add_issue(report, "errors", "STYLE_LINE_SPACING", f"{name} 的行距不是规范值 {line} twips", str(style_line_twips(style)))
    if first_line_chars is not None:
        actual_chars, actual_twips = style_indent(style)
        if actual_chars != first_line_chars and actual_twips != "560":
            add_issue(
                report,
                "errors",
                "STYLE_FIRST_INDENT",
                f"{name} 没有设置首行 {first_line_chars} 字符缩进",
                f"firstLineChars={actual_chars}, firstLine={actual_twips}",
            )
    if require_num_level is not None:
        ilvl, num_id = style_numpr(style)
        if ilvl != require_num_level or not num_id:
            add_issue(
                report,
                "errors",
                "STYLE_NUMBERING",
                f"{name} 没有绑定模板多级编号的第 {require_num_level} 级",
                f"ilvl={ilvl}, numId={num_id}",
            )
    return style


def check_heading_style_effects(report: dict, doc) -> None:
    """检查标题及其关联字符样式的最终可见属性。"""
    targets = (
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Abstract Heading",
        "English Abstract Heading",
        "TOC Title",
        "Appendix Heading",
    )
    for name in targets:
        style = find_style(doc, {name}, {name.replace(" ", "")})
        if style is None:
            continue
        issues = effective_style_decoration_issues(doc, style)
        linked = style_linked_character(doc, style)
        link_id = attribute(child(style_xml(style), "link"), "val")
        if link_id and linked is None:
            issues.append(f"关联字符样式不存在={link_id}")
        if linked is not None:
            linked_issues = effective_style_decoration_issues(doc, linked)
            if style_run_boolean(doc, linked, "b"):
                linked_issues.append("关联字符样式：不应加粗")
            expected_font = {
                "Title": ("黑体", "Times New Roman", 16),
                "Heading 1": ("黑体", "Times New Roman", 16),
                "Heading 2": ("黑体", "Times New Roman", 14),
                "Heading 3": ("黑体", "Times New Roman", 12),
            }.get(name)
            if expected_font is not None:
                linked_fonts = effective_style_fonts(doc, linked)
                if linked_fonts["eastAsia"] != expected_font[0] or linked_fonts["ascii"] != expected_font[1]:
                    linked_issues.append(f"字体={linked_fonts}")
                linked_size = style_size_pt(linked)
                if linked_size is None or abs(linked_size - expected_font[2]) > 0.01:
                    linked_issues.append(f"字号={linked_size}")
            if linked_issues:
                issues.append("关联字符样式：" + ", ".join(linked_issues))
        if issues:
            add_issue(report, "errors", "STYLE_EFFECTIVE_DECORATION", f"{name} 样式仍含非规范视觉属性", "; ".join(issues))

        if style_run_boolean(doc, style, "b"):
            add_issue(report, "errors", "STYLE_BOLD", f"{name} 不应设置加粗属性")
        if style_run_boolean(doc, style, "i"):
            add_issue(report, "errors", "STYLE_ITALIC", f"{name} 不应继承倾斜属性")

        indent = effective_style_indent(doc, style)
        nonzero_indent = [
            f"{name}={value}"
            for name, value in indent.items()
            if value not in {None, "0"}
        ]
        if nonzero_indent:
            add_issue(report, "errors", "HEADING_INDENT", f"{name} 不应有首行或左侧缩进", ", ".join(nonzero_indent))
        if effective_style_has_tabs(doc, style):
            add_issue(report, "errors", "HEADING_TABS", f"{name} 不应包含制表位定义")

        spacing = style_spacing(style)
        expected_spacing = {
            "Heading 1": {"before": "0", "after": None, "afterLines": "100", "line": "360"},
            "Heading 2": {"before": "0", "after": "0", "afterLines": None, "line": "360"},
            "Heading 3": {"before": "0", "after": "0", "afterLines": None, "line": "360"},
        }.get(name)
        if expected_spacing is not None:
            for property_name, expected in expected_spacing.items():
                if expected is not None and spacing[property_name] != expected:
                    add_issue(
                        report,
                        "errors",
                        "HEADING_SPACING",
                        f"{name} 的 {property_name} 间距不符合规范",
                        f"实际={spacing[property_name]}，期望={expected}",
                    )
                if expected is None and spacing[property_name] is not None:
                    add_issue(
                        report,
                        "errors",
                        "HEADING_SPACING",
                        f"{name} 不应设置 {property_name} 间距",
                        str(spacing[property_name]),
                    )

        ppr = style_ppr(style)
        if name == "Heading 1" and (ppr is None or child(ppr, "pageBreakBefore") is None):
            add_issue(report, "errors", "HEADING1_PAGE_BREAK", "Heading 1 没有设置章节起始分页")
        if name in {"Title", "Heading 2", "Heading 3"} and style_boolean(doc, style, "pageBreakBefore"):
            add_issue(report, "errors", "UNEXPECTED_PAGE_BREAK", f"{name} 不应设置段前分页")
        if name.startswith("Heading") and (ppr is None or child(ppr, "keepNext") is None or child(ppr, "keepLines") is None):
            add_issue(report, "warnings", "HEADING_KEEP_TOGETHER", f"{name} 未同时设置保持标题完整和与下段同页")

        if name.startswith("Heading") and not style_boolean(doc, style, "keepNext"):
            add_issue(report, "errors", "HEADING_KEEP_NEXT", f"{name} 没有设置与下段同页")
        if name.startswith("Heading") and not style_boolean(doc, style, "keepLines"):
            add_issue(report, "errors", "HEADING_KEEP_LINES", f"{name} 没有设置保持段中不分页")

    normal = find_style(doc, {"Normal"}, {"Normal"})
    if normal is not None:
        if style_run_boolean(doc, normal, "b"):
            add_issue(report, "errors", "NORMAL_BOLD", "Normal 不应默认加粗")
        if style_run_boolean(doc, normal, "i"):
            add_issue(report, "errors", "NORMAL_ITALIC", "Normal 不应默认倾斜")
        for property_name, code, message in (
            ("pageBreakBefore", "NORMAL_PAGE_BREAK", "Normal 不应设置段前分页"),
            ("keepNext", "NORMAL_KEEP_NEXT", "Normal 不应默认与下段同页"),
            ("keepLines", "NORMAL_KEEP_LINES", "Normal 不应默认保持段中不分页"),
        ):
            if style_boolean(doc, normal, property_name):
                add_issue(report, "errors", code, message)

    code = find_style(doc, {"Code Block", "CodeBlock"}, {"CodeBlock"})
    if code is not None and style_alignment(code) != WD_ALIGN_PARAGRAPH.LEFT:
        add_issue(report, "errors", "CODE_ALIGNMENT", "代码或技术文本样式没有左对齐")
    if code is not None and not style_run_boolean(doc, code, "noProof"):
        add_issue(report, "warnings", "CODE_PROOFING", "代码或技术文本样式未关闭拼写检查")


def check_style_inheritance(report: dict, doc) -> None:
    """记录关键样式继承链，发现循环或悬空继承。"""
    targets = ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3")
    report["style_inheritance"] = {}
    for name in targets:
        style = find_style(doc, {name}, {name.replace(" ", "")})
        if style is None:
            continue
        chain, cyclic, missing = style_chain(doc, style)
        report["style_inheritance"][name] = [item.style_id for item in chain]
        if cyclic:
            add_issue(report, "errors", "STYLE_INHERITANCE_CYCLE", f"{name} 存在循环继承", " -> ".join(item.style_id for item in chain))
        if missing:
            add_issue(report, "errors", "STYLE_INHERITANCE_MISSING", f"{name} 的继承样式不存在", missing)


def numbering_level(xml_parts: dict[str, object], num_id: str | None, ilvl: str | None):
    if not num_id or not ilvl:
        return None, None
    root = xml_parts.get("word/numbering.xml")
    if root is None:
        return None, None
    abstract_id = None
    for num in root.findall(W + "num"):
        if attribute(num, "numId") == num_id:
            abstract_id = attribute(child(num, "abstractNumId"), "val")
            break
    if abstract_id is None:
        return None, None
    for abstract in root.findall(W + "abstractNum"):
        if attribute(abstract, "abstractNumId") != abstract_id:
            continue
        for level in abstract.findall(W + "lvl"):
            if attribute(level, "ilvl") == ilvl:
                return abstract, level
    return None, None


def check_numbering_bindings(report: dict, doc, xml_parts: dict[str, object]) -> None:
    """解析 numId 到 abstractNum，防止只看到编号属性却忽略实际编号格式。"""
    expected = {
        "Heading 1": ("0", "chineseCounting", "第%1章", False, "center"),
        "Heading 2": ("1", "decimal", "%1.%2", True, "left"),
        "Heading 3": ("2", "decimal", "%1.%2.%3", True, "left"),
    }
    for name, (expected_level, expected_format, expected_text, expected_legal, expected_alignment) in expected.items():
        style = find_style(doc, {name}, {name.replace(" ", "")})
        if style is None:
            continue
        actual_level, num_id = style_numpr(style)
        abstract, level = numbering_level(xml_parts, num_id, actual_level)
        if abstract is None or level is None:
            add_issue(report, "errors", "NUMBERING_CHAIN", f"{name} 的编号链无法解析", f"ilvl={actual_level}, numId={num_id}")
            continue
        actual_format = attribute(child(level, "numFmt"), "val")
        actual_text = attribute(child(level, "lvlText"), "val")
        actual_alignment = attribute(child(level, "lvlJc"), "val")
        level_style = attribute(child(child(level, "pPr"), "pStyle"), "val")
        actual_legal = child(level, "isLgl") is not None
        if actual_level != expected_level or actual_format != expected_format or actual_text != expected_text:
            add_issue(
                report,
                "errors",
                "NUMBERING_DEFINITION",
                f"{name} 的实际编号格式不符合规范",
                f"级别={actual_level}, 格式={actual_format}, 文本={actual_text}",
            )
        if actual_alignment != expected_alignment:
            add_issue(
                report,
                "errors",
                "NUMBERING_ALIGNMENT",
                f"{name} 的编号级别没有左/中对齐到规范位置",
                f"实际={actual_alignment}，期望={expected_alignment}",
            )
        indent = child(child(level, "pPr"), "ind")
        actual_indent = {
            name: attribute(indent, name)
            for name in ("left", "hanging", "firstLine")
        }
        if any(value != "0" for value in actual_indent.values()):
            add_issue(
                report,
                "errors",
                "NUMBERING_INDENT",
                f"{name} 的编号级别含有隐式缩进或制表位前置空间",
                str(actual_indent),
            )
        if actual_legal != expected_legal:
            add_issue(
                report,
                "errors",
                "NUMBERING_LEGAL_MODE",
                f"{name} 的编号没有正确设置 Word 阿拉伯数字兼容模式",
                f"isLgl={actual_legal}，期望={expected_legal}",
            )
        if level_style not in {style.style_id, name.replace(" ", "")}:
            add_issue(report, "errors", "NUMBERING_STYLE_LINK", f"{name} 的编号级别没有绑定自身样式", str(level_style))


def load_xml(zfile: ZipFile, name: str):
    try:
        return ET.fromstring(zfile.read(name))
    except KeyError:
        return None


def xml_text(root) -> str:
    if root is None:
        return ""
    return "".join(node.text or "" for node in root.iter(W + "t"))


def all_field_instructions(root) -> str:
    if root is None:
        return ""
    values = [node.text or "" for node in root.iter(W + "instrText")]
    values.extend(node.get(W + "instr") or "" for node in root.iter(W + "fldSimple"))
    return " ".join(values)


def detect_profile(doc) -> tuple[str, list[str]]:
    texts = [clean(paragraph.text) for paragraph in doc.paragraphs]
    signals: list[str] = []
    if any(text in {"摘要", "摘 要"} for text in texts):
        signals.append("中文摘要")
    if any(text == "ABSTRACT" for text in texts):
        signals.append("英文摘要")
    if any(text in {"目录", "目 录"} for text in texts):
        signals.append("目录")
    if any(text == "参考文献" for text in texts):
        signals.append("参考文献")
    if any(re.match(r"^第[一二三四五六七八九十百零]+章", text) for text in texts):
        signals.append("章标题")
    if len(signals) >= 2:
        return "thesis", signals
    return "general", signals


def check_baseline_marker(report: dict, xml_parts: dict[str, object], require_baseline: bool) -> None:
    settings = xml_parts.get("word/settings.xml")
    marker_found = False
    if settings is not None:
        for variable in settings.iter(W + "docVar"):
            if attribute(variable, "name") == BASELINE_MARKER_NAME and attribute(variable, "val") == BASELINE_MARKER_VALUE:
                marker_found = True
                break
    report["baseline"] = {"required": require_baseline, "verified": marker_found}
    if marker_found:
        record_check(report, "BASELINE_MARKER", True, "发现中文 DOCX 基线模板标识")
    elif require_baseline:
        add_issue(report, "errors", "BASELINE_MISSING", "没有找到本技能空白基线模板标识")
    else:
        add_issue(report, "warnings", "BASELINE_UNVERIFIED", "未验证文档是否来源于本技能空白基线模板")


def check_legacy_styles(report: dict, doc) -> None:
    """发现源文档留下的论文专用样式，但不把合法的用户正文当作错误。"""
    markers = ("lyh_", "mathtype", "endnote bibliography", "mtdisplayequation")
    residue = [
        f"{style.style_id}:{style.name}"
        for style in doc.styles
        if any(marker in (style.style_id + " " + style.name).lower() for marker in markers)
    ]
    report["legacy_styles"] = residue
    if residue:
        add_issue(report, "warnings", "LEGACY_STYLE_RESIDUE", "发现源文档专用样式残留", "; ".join(residue[:10]))
    else:
        record_check(report, "LEGACY_STYLES", True, "没有发现论文源文档专用样式残留")


def check_page_and_sections(report: dict, doc, xml_parts: dict[str, object], profile: str) -> None:
    expected = (21.0, 29.7, 2.0, 2.0, 2.0, 2.7)
    for index, section in enumerate(doc.sections, start=1):
        width = float(section.page_width) / EMU_PER_INCH * CM_PER_INCH
        height = float(section.page_height) / EMU_PER_INCH * CM_PER_INCH
        margins = tuple(
            float(value) / EMU_PER_INCH * CM_PER_INCH
            for value in (
                section.top_margin,
                section.left_margin,
                section.right_margin,
                section.bottom_margin,
            )
        )
        values = (width, height, *margins)
        passed = all(abs(actual - wanted) <= TOLERANCE_CM for actual, wanted in zip(values, expected))
        record_check(report, f"SECTION_{index}", passed, f"第 {index} 节页面和边距")
        if not passed:
            add_issue(
                report,
                "errors",
                "SECTION_LAYOUT",
                f"第 {index} 节不是 A4 或边距不符合规范",
                "实际值 cm：宽 %.2f，高 %.2f，上 %.2f，左 %.2f，右 %.2f，下 %.2f"
                % values,
            )

        if profile == "thesis":
            header_distance = float(section.header_distance) / EMU_PER_INCH * CM_PER_INCH
            footer_distance = float(section.footer_distance) / EMU_PER_INCH * CM_PER_INCH
            distance_passed = (
                abs(header_distance - 1.5) <= TOLERANCE_CM
                and abs(footer_distance - 1.75) <= TOLERANCE_CM
            )
            record_check(report, f"SECTION_DISTANCE_{index}", distance_passed, f"第 {index} 节页眉页脚距离")
            if not distance_passed:
                add_issue(
                    report,
                    "errors",
                    "SECTION_HEADER_FOOTER_DISTANCE",
                    f"第 {index} 节页眉/页脚距离不符合论文模式",
                    "实际页眉 %.2f cm，页脚 %.2f cm；期望页眉 1.50 cm，页脚 1.75 cm"
                    % (header_distance, footer_distance),
                )
            sect_pr = doc.sections[index - 1]._sectPr
            if sect_pr.find(W + "titlePg") is None:
                add_issue(report, "errors", "THESIS_FIRST_PAGE_HEADER", f"第 {index} 节未设置论文首页页眉分隔")

    document_xml = xml_parts.get("word/document.xml")
    starts = [
        attribute(node, "start")
        for node in document_xml.iter(W + "pgNumType")
    ] if document_xml is not None else []
    if "1" not in starts:
        add_issue(report, "errors", "PAGE_START", "没有找到正文从 1 开始的页码节属性")
    else:
        record_check(report, "PAGE_START", True, "正文页码起点为 1")

    roman_formats = [
        attribute(node, "fmt")
        for node in (document_xml.iter(W + "pgNumType") if document_xml is not None else [])
        if attribute(node, "fmt") in {"upperRoman", "lowerRoman"}
    ]
    if roman_formats:
        add_issue(report, "errors", "PAGE_ROMAN_FORMAT", "页码不应使用前置罗马数字模式", str(roman_formats))
    else:
        record_check(report, "PAGE_ARABIC_FORMAT", True, "未发现前置罗马数字页码定义")

    settings = xml_parts.get("word/settings.xml")
    even_odd = settings.find(W + "evenAndOddHeaders") if settings is not None else None
    if settings is not None and list(settings.iter(W + "docGrid")):
        add_issue(report, "errors", "DOCUMENT_GRID_PRESENT", "文档不应启用文档网格")
    character_spacing = list(settings.iter(W + "characterSpacingControl")) if settings is not None else []
    if character_spacing:
        add_issue(report, "errors", "PUNCTUATION_COMPRESSION", "文档不应启用字符间距或标点压缩设置")
    if profile == "thesis" and even_odd is None:
        add_issue(report, "errors", "THESIS_ODD_EVEN_HEADERS", "论文模式未启用奇偶页眉设置")
    elif profile == "thesis":
        record_check(report, "THESIS_ODD_EVEN_HEADERS", True, "论文模式已启用奇偶页眉设置")

    footer_roots = [root for name, root in xml_parts.items() if name.startswith("word/footer") and root is not None]
    page_field_found = False
    for root in footer_roots:
        for paragraph in root.iter(W + "p"):
            instructions = all_field_instructions(paragraph)
            if "PAGE" not in instructions.upper():
                continue
            page_field_found = True
            alignment = paragraph.find(W + "pPr/" + W + "jc")
            direct_alignment = attribute(alignment, "val")
            style_node = paragraph.find(W + "pPr/" + W + "pStyle")
            style = find_style(doc, set(), {attribute(style_node, "val")} if style_node is not None else set())
            effective_alignment = direct_alignment or (
                {WD_ALIGN_PARAGRAPH.LEFT: "left", WD_ALIGN_PARAGRAPH.CENTER: "center", WD_ALIGN_PARAGRAPH.RIGHT: "right", WD_ALIGN_PARAGRAPH.JUSTIFY: "both"}.get(style_alignment(style))
                if style is not None
                else None
            )
            if effective_alignment != "center":
                add_issue(report, "errors", "PAGE_ALIGNMENT", "页码字段所在段落没有居中")
    if not page_field_found:
        add_issue(report, "errors", "PAGE_FIELD", "页脚没有找到 PAGE 页码字段")
    else:
        record_check(report, "PAGE_FIELD", True, "找到页脚 PAGE 字段")

    header_text = " ".join(
        xml_text(root)
        for name, root in xml_parts.items()
        if name.startswith("word/header") and root is not None
    ).strip()
    if header_text and profile != "thesis":
        add_issue(report, "warnings", "HEADER_PRESENT", "发现非空页眉；默认基线不设置固定页眉")
    elif profile == "thesis":
        record_check(report, "THESIS_HEADER", bool(header_text), "论文模式保留用户提供的页眉文字；未虚构学校信息")


def check_styles(report: dict, doc, profile: str) -> None:
    common = [
        ("Normal", {"Normal"}, "宋体", "Times New Roman", 12, WD_ALIGN_PARAGRAPH.JUSTIFY, "360", "200", None),
        ("Title", {"Title"}, "黑体", "Times New Roman", 16, WD_ALIGN_PARAGRAPH.CENTER, "360", None, None),
        ("Heading 1", {"Heading1"}, "黑体", "Times New Roman", 16, WD_ALIGN_PARAGRAPH.CENTER, "360", None, "0"),
        ("Heading 2", {"Heading2"}, "黑体", "Times New Roman", 14, WD_ALIGN_PARAGRAPH.LEFT, "360", None, "1"),
        ("Heading 3", {"Heading3"}, "黑体", "Times New Roman", 12, WD_ALIGN_PARAGRAPH.LEFT, "360", None, "2"),
        ("Figure Caption", {"Figure Caption"}, "楷体", "Times New Roman", 10.5, WD_ALIGN_PARAGRAPH.CENTER, "300", None, None),
        ("Table Caption", {"Table Caption"}, "楷体", "Times New Roman", 10.5, WD_ALIGN_PARAGRAPH.CENTER, "300", None, None),
        ("Figure List", {"FigureList"}, "楷体", "Times New Roman", 10.5, WD_ALIGN_PARAGRAPH.LEFT, "300", None, None),
        ("Table List", {"TableList"}, "楷体", "Times New Roman", 10.5, WD_ALIGN_PARAGRAPH.LEFT, "300", None, None),
        ("Equation", {"Equation"}, "Cambria Math", "Cambria Math", 12, WD_ALIGN_PARAGRAPH.CENTER, None, None, None),
        ("Reference", {"Reference"}, "宋体", "Times New Roman", 12, None, None, None, None),
        ("Abstract Body", {"Abstract Body"}, "宋体", "Times New Roman", 14, WD_ALIGN_PARAGRAPH.JUSTIFY, "360", "200", None),
        ("English Abstract Body", {"English Abstract Body"}, "Times New Roman", "Times New Roman", 14, WD_ALIGN_PARAGRAPH.JUSTIFY, "360", "400", None),
        ("Abstract Heading", {"Abstract Heading"}, "黑体", "Times New Roman", 16, WD_ALIGN_PARAGRAPH.CENTER, "360", None, None),
        ("English Abstract Heading", {"English Abstract Heading"}, "黑体", "Times New Roman", 16, WD_ALIGN_PARAGRAPH.CENTER, "360", None, None),
        ("TOC Title", {"TOC Title"}, "黑体", "Times New Roman", 16, WD_ALIGN_PARAGRAPH.CENTER, "360", None, None),
    ]
    for name, ids, east_asia, latin, size, alignment, line, first_line, num_level in common:
        check_style(
            report,
            doc,
            name,
            ids,
            east_asia=east_asia,
            latin=latin,
            size=size,
            alignment=alignment,
            line=line,
            first_line_chars=first_line,
            require_num_level=num_level,
        )

    reference = find_style(doc, {"Reference"}, {"Reference"})
    if reference is not None:
        ppr = style_ppr(reference)
        indent = ppr.find(W + "ind") if ppr is not None else None
        if attribute(indent, "hanging") != "480":
            add_issue(report, "errors", "REFERENCE_INDENT", "Reference 样式没有约定的悬挂缩进")

    for level in range(1, 4):
        style = find_style(doc, {f"Heading {level}"}, {f"Heading{level}"})
        if style is not None:
            record_check(report, f"HEADING_STYLE_{level}", True, f"Heading {level} 样式存在")

    for level in range(1, 4):
        check_style(
            report,
            doc,
            f"TOC {level}",
            {f"TOC{level}"},
            east_asia="宋体",
            latin="Times New Roman",
            size=12,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            line="300",
        )

    check_style_inheritance(report, doc)
    check_heading_style_effects(report, doc)


def check_direct_formatting(report: dict, doc) -> None:
    direct_paragraphs = 0
    direct_runs = 0
    table_direct_paragraphs = 0
    allowed_ppr = {W + "pStyle", W + "numPr", W + "sectPr"}
    allowed_rpr = {W + "rStyle", W + "vertAlign", W + "lang", W + "rtl", W + "noProof"}
    for paragraph in all_paragraphs(doc):
        is_page_field = paragraph_has_field(paragraph, "PAGE")
        ppr = paragraph._p.find(W + "pPr")
        if ppr is not None:
            disallowed = [item.tag for item in ppr if item.tag not in allowed_ppr]
            if paragraph_in_table(paragraph) and disallowed == [W + "jc"]:
                table_direct_paragraphs += 1
            elif disallowed and not is_page_field:
                direct_paragraphs += 1
        for run in paragraph.runs:
            rpr = run._r.find(W + "rPr")
            if rpr is not None and any(child.tag not in allowed_rpr for child in rpr) and not is_page_field:
                direct_runs += 1
    report["direct_formatting"] = {
        "paragraphs": direct_paragraphs,
        "table_alignment_paragraphs": table_direct_paragraphs,
        "runs": direct_runs,
    }
    if direct_paragraphs or direct_runs:
        add_issue(
            report,
            "warnings",
            "DIRECT_FORMATTING",
            "发现直接段落或字符格式；应优先回收到样式",
            f"段落 {direct_paragraphs} 处，字符 {direct_runs} 处",
        )
    else:
        record_check(report, "DIRECT_FORMATTING", True, "没有发现直接格式覆盖")


def check_headings_and_title(report: dict, doc) -> None:
    title = find_style(doc, {"Title"}, {"Title"})
    title_paragraphs = [p for p in doc.paragraphs if paragraph_style_role(p) == "Title"]
    if not title_paragraphs and any(clean(p.text) for p in doc.paragraphs):
        add_issue(report, "warnings", "TITLE_MISSING", "没有找到使用 Title 样式的文档标题")
    for paragraph in title_paragraphs:
        line_breaks = len(paragraph._p.findall(".//" + W + "br"))
        if line_breaks > 1:
            add_issue(report, "errors", "TITLE_LINES", "文档标题超过两行")
        if title is not None and style_alignment(title) != WD_ALIGN_PARAGRAPH.CENTER:
            add_issue(report, "errors", "TITLE_ALIGNMENT", "Title 样式没有居中")
        ppr = paragraph._p.find(W + "pPr")
        if ppr is not None:
            borders = ppr.find(W + "pBdr")
            if borders is not None and any(attribute(item, "val") not in {None, "none"} for item in list(borders)):
                add_issue(report, "errors", "TITLE_DIRECT_BORDER", "文档标题段落含有直接段落边框")

    previous_level = 0
    seen_heading = False
    for paragraph in doc.paragraphs:
        role = paragraph_style_role(paragraph)
        level = {"Heading 1": 1, "Heading 2": 2, "Heading 3": 3}.get(role)
        if level is None:
            continue
        seen_heading = True
        if paragraph.text != paragraph.text.lstrip(" \t\u00a0"):
            add_issue(report, "errors", "HEADING_TEXT_INDENT", f"{role} 文本开头含有空格或制表符", repr(paragraph.text[:20]))
        expected_prefix = {
            1: r"^(第[一二三四五六七八九十百零]+章|附录\s*[0-9A-Za-z]+)",
            2: r"^\d+\.\d+",
            3: r"^\d+\.\d+\.\d+",
        }[level]
        if re.match(expected_prefix, clean(paragraph.text)):
            add_issue(
                report,
                "errors",
                "MANUAL_HEADING_NUMBER",
                f"{role} 的文本包含手工编号，应只保留标题文字",
                clean(paragraph.text),
            )
        style = paragraph.style
        ilvl, num_id = style_numpr(style)
        if ilvl != str(level - 1) or not num_id:
            add_issue(report, "errors", "HEADING_NUMBERING", f"{role} 没有使用模板多级编号")
        if seen_heading and level > previous_level + 1 and previous_level:
            add_issue(report, "errors", "HEADING_LEVEL_JUMP", f"标题层级从 {previous_level} 级跳到 {level} 级")
        previous_level = level


def check_appendix(report: dict, doc, xml_parts: dict[str, object]) -> None:
    appendix_paragraphs = [
        paragraph
        for paragraph in doc.paragraphs
        if paragraph_style_role(paragraph) == "Appendix Heading"
    ]
    for paragraph in appendix_paragraphs:
        ilvl, num_id = style_numpr(paragraph.style)
        if ilvl != "0" or not num_id:
            add_issue(report, "errors", "APPENDIX_NUMBERING", "附录标题没有绑定独立的附录编号")
        abstract, level = numbering_level(xml_parts, num_id, ilvl)
        if abstract is None or level is None:
            add_issue(report, "errors", "APPENDIX_NUMBERING_CHAIN", "附录标题的编号链无法解析")
        else:
            actual_text = attribute(child(level, "lvlText"), "val")
            level_style = attribute(child(child(level, "pPr"), "pStyle"), "val")
            if actual_text != "附录%1":
                add_issue(report, "errors", "APPENDIX_NUMBERING_DEFINITION", "附录编号格式不是“附录%1”", str(actual_text))
            if level_style not in {paragraph.style.style_id, "AppendixHeading"}:
                add_issue(report, "errors", "APPENDIX_NUMBERING_STYLE_LINK", "附录编号没有绑定 Appendix Heading 样式", str(level_style))
        if re.match(r"^附录\s*[0-9A-Za-z]+", clean(paragraph.text)):
            add_issue(report, "errors", "MANUAL_APPENDIX_NUMBER", "附录标题文本包含手工编号", clean(paragraph.text))
    if appendix_paragraphs:
        record_check(report, "APPENDIX", True, f"检查 {len(appendix_paragraphs)} 个附录标题")


def keyword_parts(text: str, label: str, require_separator: bool = False) -> list[str]:
    separator = r"\s*[：:]\s*" if require_separator else r"\s*(?:[：:]\s*)?"
    match = re.search(rf"^{re.escape(label)}{separator}(.*)$", text, re.IGNORECASE)
    if not match:
        return []
    value = match.group(1).strip()
    value = re.sub(r"[。；;，,]+$", "", value).strip()
    return [part.strip() for part in re.split(r"[；;]", value) if part.strip()]


def check_abstracts(report: dict, doc) -> None:
    texts = [clean(p.text) for p in doc.paragraphs]
    chinese_index = next((i for i, text in enumerate(texts) if text in {"摘要", "摘 要"}), None)
    if chinese_index is not None:
        abstract_style = find_style(doc, {"Abstract Body"}, {"Abstract Body"})
        keyword_index = next(
            (i for i in range(chinese_index + 1, len(texts)) if texts[i].startswith("关键词")),
            None,
        )
        if keyword_index is None:
            add_issue(report, "errors", "KEYWORDS_MISSING", "有中文摘要但没有关键词行")
            keyword_index = len(texts)
        body = [
            doc.paragraphs[i]
            for i in range(chinese_index + 1, keyword_index)
            if texts[i]
        ]
        if not body:
            add_issue(report, "errors", "ABSTRACT_BODY_MISSING", "中文摘要标题后没有摘要正文")
        for paragraph in body:
            if paragraph_style_role(paragraph) != "Abstract Body":
                add_issue(report, "errors", "ABSTRACT_STYLE", "中文摘要正文没有使用 Abstract Body 样式", clean(paragraph.text))
        if keyword_index < len(texts):
            parts = keyword_parts(texts[keyword_index], "关键词", require_separator=True)
            if not 4 <= len(parts) <= 6:
                add_issue(report, "errors", "KEYWORD_COUNT", "中文关键词数量必须为 4 至 6 个", str(len(parts)))
            if not re.match(r"^关键词\s*[：:]", texts[keyword_index]):
                add_issue(report, "errors", "KEYWORD_SEPARATOR", "中文关键词标签后必须有中文冒号")
            if texts[keyword_index].rstrip().endswith(("。", "；", ";", "，", ",")):
                add_issue(report, "errors", "KEYWORD_ENDING", "最后一个中文关键词后不得有标点")
        record_check(report, "CHINESE_ABSTRACT", True, "发现并检查中文摘要模块")

    english_index = next((i for i, text in enumerate(texts) if text == "ABSTRACT"), None)
    if english_index is not None:
        keywords_index = next(
            (i for i in range(english_index + 1, len(texts)) if texts[i].upper().startswith("KEYWORDS")),
            None,
        )
        if keywords_index is None:
            add_issue(report, "errors", "EN_KEYWORDS_MISSING", "有英文摘要但没有 KEYWORDS 行")
            keywords_index = len(texts)
        body = [
            doc.paragraphs[i]
            for i in range(english_index + 1, keywords_index)
            if texts[i]
        ]
        if not body:
            add_issue(report, "errors", "EN_ABSTRACT_BODY_MISSING", "ABSTRACT 后没有英文摘要正文")
        for paragraph in body:
            if paragraph_style_role(paragraph) != "English Abstract Body":
                add_issue(report, "errors", "EN_ABSTRACT_STYLE", "英文摘要正文没有使用 English Abstract Body 样式", clean(paragraph.text))
        if keywords_index < len(texts):
            parts = keyword_parts(texts[keywords_index], "KEYWORDS")
            if not parts:
                add_issue(report, "errors", "EN_KEYWORDS_EMPTY", "KEYWORDS 后没有关键词")
            for part in parts:
                if any(char.isalpha() and char.isascii() and char != char.lower() for char in part):
                    add_issue(report, "errors", "EN_KEYWORDS_CASE", "英文关键词必须使用小写", part)
            if texts[keywords_index].rstrip().endswith(("。", "；", ";", "，", ",")):
                add_issue(report, "errors", "EN_KEYWORDS_ENDING", "最后一个英文关键词后不得有标点")
        record_check(report, "ENGLISH_ABSTRACT", True, "发现并检查英文摘要模块")


def check_toc(report: dict, doc, xml_parts: dict[str, object]) -> None:
    texts = [clean(p.text) for p in doc.paragraphs]
    toc_root = xml_parts.get("word/document.xml")
    instructions = all_field_instructions(toc_root)
    has_toc_title = any(text in {"目录", "目 录"} for text in texts)
    has_toc_field = "TOC" in instructions.upper()
    cached_entries = []
    if toc_root is not None:
        for paragraph in toc_root.iter(W + "p"):
            if not any(node.tag == W + "hyperlink" for node in paragraph.iter()):
                continue
            value = clean(xml_text(paragraph))
            if value:
                cached_entries.append(value)
    report["toc"] = {
        "field_present": has_toc_field,
        "update_fields": None,
        "cached_entry_count": len(cached_entries),
        "cached_entries": cached_entries[:20],
    }
    if has_toc_title and not has_toc_field:
        add_issue(report, "errors", "TOC_FIELD_MISSING", "存在目录标题但没有真实 TOC 字段")
    if has_toc_field:
        if not re.search(r"\\o\s+[\"“]?1-3", instructions, re.IGNORECASE):
            add_issue(report, "warnings", "TOC_LEVEL_RANGE", "TOC 字段未明确覆盖 Heading 1 至 Heading 3", instructions)
        settings = xml_parts.get("word/settings.xml")
        update_value = None
        if settings is not None:
            update = settings.find(W + "updateFields")
            update_value = attribute(update, "val")
        report["toc"]["update_fields"] = update_value
        if update_value != "true":
            add_issue(report, "warnings", "TOC_UPDATE_FIELDS", "TOC 字段未设置打开文档时更新")
        if not cached_entries:
            add_issue(report, "warnings", "TOC_CACHE_EMPTY", "TOC 字段没有缓存条目；需在 Word 中更新目录")
        mixed = [value for value in cached_entries if re.match(r"^[一二三四五六七八九十百千万零〇]+[.．]", value)]
        if mixed:
            add_issue(report, "errors", "TOC_MIXED_NUMBERING", "目录缓存含有中文数字与阿拉伯小节号混排", str(mixed[:5]))
        record_check(report, "TOC_FIELD", True, "找到真实 TOC 字段")


def caption_number(text: str, label: str) -> tuple[int, int] | None:
    match = re.match(rf"^{re.escape(label)}\s*([0-9]+)\.([0-9]+)", text)
    if not match:
        appendix_match = re.match(rf"^{re.escape(label)}\s*[A-Za-z]+([0-9]+)", text)
        if appendix_match:
            return -1, int(appendix_match.group(1))
        return None
    return int(match.group(1)), int(match.group(2))


def section_text_width_twips(section) -> int:
    page_width = int(round(float(section.page_width) / EMU_PER_INCH * TWIPS_PER_INCH))
    left = int(round(float(section.left_margin) / EMU_PER_INCH * TWIPS_PER_INCH))
    right = int(round(float(section.right_margin) / EMU_PER_INCH * TWIPS_PER_INCH))
    return page_width - left - right


def table_widths(table) -> tuple[int | None, list[int]]:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    width_node = tbl_pr.find(W + "tblW") if tbl_pr is not None else None
    width = None
    if width_node is not None and attribute(width_node, "type") in {None, "dxa"}:
        value = attribute(width_node, "w")
        if value and value.lstrip("-").isdigit():
            width = int(value)
    grid = tbl.find(W + "tblGrid")
    grid_widths: list[int] = []
    if grid is not None:
        for column in grid.findall(W + "gridCol"):
            value = attribute(column, "w")
            if value and value.lstrip("-").isdigit():
                grid_widths.append(int(value))
    return width, grid_widths


def table_border_issues(table) -> list[str]:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(W + "tblBorders") if tbl_pr is not None else None
    if borders is None:
        return ["未设置明确网格边框"]
    required = ("top", "left", "bottom", "right", "insideH", "insideV")
    issues = []
    for side in required:
        node = borders.find(W + side)
        if node is None or attribute(node, "val") in {None, "none"}:
            issues.append(side)
    return issues


def first_row_repeats(table) -> bool:
    rows = table._tbl.findall(W + "tr")
    if not rows:
        return False
    tr_pr = rows[0].find(W + "trPr")
    return tr_pr is not None and tr_pr.find(W + "tblHeader") is not None


def check_captions_tables_images(report: dict, doc, xml_parts: dict[str, object]) -> None:
    figure_numbers: list[tuple[int, int]] = []
    table_numbers: list[tuple[int, int]] = []
    for paragraph in doc.paragraphs:
        role = paragraph_style_role(paragraph)
        text = clean(paragraph.text)
        if role in {"Figure Caption", "Table Caption"} and not text:
            continue
        if role == "Figure Caption":
            number = caption_number(text, "图")
            if number is None:
                add_issue(report, "errors", "FIGURE_CAPTION_NUMBER", "图题没有使用章号.图号格式", text)
            else:
                figure_numbers.append(number)
        elif role == "Table Caption":
            number = caption_number(text, "表")
            if number is None:
                add_issue(report, "errors", "TABLE_CAPTION_NUMBER", "表题没有使用章号.表号格式", text)
            else:
                table_numbers.append(number)
        elif re.match(r"^(图|表)\s*[0-9]", text) and not (role and (role.startswith("TOC") or role.endswith("List"))) and role not in {"Figure Caption", "Table Caption"}:
            add_issue(report, "errors", "CAPTION_STYLE", "图题或表题没有使用规范样式", text)

    for label, numbers in (("图", figure_numbers), ("表", table_numbers)):
        grouped: defaultdict[int, list[int]] = defaultdict(list)
        for chapter, item in numbers:
            grouped[chapter].append(item)
        for chapter, items in grouped.items():
            if chapter == -1:
                continue
            expected = list(range(1, len(items) + 1))
            if sorted(items) != expected:
                add_issue(
                    report,
                    "errors",
                    "CAPTION_SEQUENCE",
                    f"{label} {chapter} 章内编号不连续",
                    str(sorted(items)),
                )

    document_xml = xml_parts.get("word/document.xml")
    anchors = len(list(document_xml.iter(WP + "anchor"))) if document_xml is not None else 0
    inline = len(list(document_xml.iter(WP + "inline"))) if document_xml is not None else 0
    report["images"] = {"inline": inline, "floating": anchors}
    if anchors:
        add_issue(report, "warnings", "FLOATING_IMAGE", f"发现 {anchors} 个浮动图像对象，默认应优先使用行内对象")

    available_widths = [section_text_width_twips(section) for section in doc.sections]
    available_width = min(available_widths) if available_widths else None
    for table in all_tables(doc):
        tbl = table._tbl
        tbl_pr = tbl.tblPr
        style_node = tbl_pr.find(W + "tblStyle") if tbl_pr is not None else None
        style_value = attribute(style_node, "val")
        borders = tbl_pr.find(W + "tblBorders") if tbl_pr is not None else None
        grid_style = style_value and ("grid" in style_value.lower() or "网格" in style_value)
        border_issues = table_border_issues(table)
        if border_issues and not grid_style:
            add_issue(report, "errors", "TABLE_BORDERS", "表格没有完整的外框和内部网格边框", ", ".join(border_issues))
        elif border_issues:
            add_issue(report, "warnings", "TABLE_STYLE_BORDERS", "表格依赖网格样式而非完整显式边框", ", ".join(border_issues))

        table_width, grid_widths = table_widths(table)
        if table_width is None:
            add_issue(report, "errors", "TABLE_WIDTH_MISSING", "表格没有设置明确的总宽度")
        if not grid_widths:
            add_issue(report, "errors", "TABLE_GRID_MISSING", "表格没有设置明确的列宽网格")
        if available_width is not None:
            if table_width is not None and table_width > available_width + 24:
                add_issue(report, "errors", "TABLE_WIDTH", "表格总宽度超过版心", f"表格={table_width} twips，版心={available_width} twips")
            if grid_widths and sum(grid_widths) > available_width + 24:
                add_issue(report, "errors", "TABLE_GRID_WIDTH", "表格列宽总和超过版心", f"列宽总和={sum(grid_widths)} twips，版心={available_width} twips")
        if table_width is not None and grid_widths and abs(sum(grid_widths) - table_width) > 24:
            add_issue(
                report,
                "errors",
                "TABLE_GRID_MISMATCH",
                "表格总宽度与列宽网格不一致",
                f"表格={table_width} twips，列宽总和={sum(grid_widths)} twips",
            )
        if len(tbl.findall(W + "tr")) >= 6 and not first_row_repeats(table):
            add_issue(report, "warnings", "TABLE_HEADER_REPEAT", "长表格首行没有设置跨页重复")
        for row in tbl.findall(W + "tr"):
            tr_pr = row.find(W + "trPr")
            height = tr_pr.find(W + "trHeight") if tr_pr is not None else None
            if height is not None and attribute(height, "hRule") == "exact":
                add_issue(report, "errors", "TABLE_FIXED_HEIGHT", "表格使用 exact 固定行高，存在截断风险")


def check_equations(report: dict, xml_parts: dict[str, object]) -> None:
    document_xml = xml_parts.get("word/document.xml")
    if document_xml is None:
        return
    omath = list(document_xml.iter(M + "oMath"))
    omath_para = list(document_xml.iter(M + "oMathPara"))
    embedding_names = [
        name for name in xml_parts["zip_names"]
        if name.startswith("word/embeddings/")
    ]
    math_type_names = [name for name in embedding_names if "mathtype" in name.lower()]
    if embedding_names:
        add_issue(
            report,
            "errors",
            "OLE_EMBEDDING",
            "文档含有嵌入 OLE 对象；本技能的公式目标格式为 OMML",
            ", ".join(embedding_names[:5]),
        )
    if omath:
        if not omath_para:
            add_issue(report, "warnings", "OMML_INLINE", "发现 OMML，但没有 m:oMathPara；请确认公式是否另起一行")
        instructions = all_field_instructions(document_xml)
        if not re.search(r"SEQ\s+(Equation|公式)", instructions, re.IGNORECASE):
            add_issue(report, "errors", "EQUATION_NUMBERING", "发现 OMML 公式但没有 Equation 或公式序号字段")
        record_check(report, "OMML", True, f"发现 {len(omath)} 个可编辑 OMML 公式")
    elif math_type_names:
        add_issue(report, "errors", "MATHTYPE_OBJECT", "文档仍含 MathType 对象", ", ".join(math_type_names[:5]))

    raw_latex = re.findall(r"\\(?:frac|sqrt|sum|int|alpha|beta|begin|end)\b", xml_text(document_xml))
    if raw_latex:
        add_issue(report, "errors", "RAW_LATEX", "文档文本中残留原始 LaTeX", ", ".join(sorted(set(raw_latex))))


def reference_blocks(doc) -> list[tuple[int, object]]:
    heading_index = next(
        (i for i, paragraph in enumerate(doc.paragraphs) if clean(paragraph.text) == "参考文献"),
        None,
    )
    if heading_index is None:
        return []
    result = []
    for index in range(heading_index + 1, len(doc.paragraphs)):
        paragraph = doc.paragraphs[index]
        text = clean(paragraph.text)
        role = paragraph_style_role(paragraph)
        if not text:
            continue
        if index > heading_index + 1 and role in {"Heading 1", "Title"} and not re.match(r"^[［\[]\s*\d+", text):
            break
        if re.match(r"^[［\[]\s*\d+", text) or role == "Reference":
            result.append((index, paragraph))
        elif result:
            break
    return result


def check_references_and_citations(report: dict, doc) -> None:
    blocks = reference_blocks(doc)
    if not blocks:
        if any(clean(p.text) == "参考文献" for p in doc.paragraphs):
            add_issue(report, "errors", "REFERENCE_EMPTY", "存在参考文献标题但没有参考文献条目")
        return

    expected = 1
    reference_numbers: set[int] = set()
    for _, paragraph in blocks:
        text = clean(paragraph.text)
        match = re.match(r"^[［\[]\s*(\d+)\s*[］\]]", text)
        if not match:
            add_issue(report, "errors", "REFERENCE_PREFIX", "参考文献条目缺少方括号数字序号", text)
            continue
        number = int(match.group(1))
        reference_numbers.add(number)
        if number != expected:
            add_issue(report, "errors", "REFERENCE_SEQUENCE", "参考文献序号不连续或未从 1 开始", f"期望 {expected}，实际 {number}")
            expected = number + 1
        else:
            expected += 1
        if paragraph_style_role(paragraph) != "Reference":
            add_issue(report, "errors", "REFERENCE_STYLE", "参考文献条目没有使用 Reference 样式", text)
        if text.startswith("["):
            add_issue(report, "warnings", "ASCII_REFERENCE_BRACKET", "参考文献默认应使用全角方括号", text[:30])

    citation_numbers: set[int] = set()
    for paragraph in doc.paragraphs:
        if paragraph in [item[1] for item in blocks]:
            continue
        for match in re.finditer(r"［(\d+)］|\[(\d+)\]", paragraph.text):
            citation_numbers.add(int(match.group(1) or match.group(2)))
            if match.group(0).startswith("["):
                add_issue(report, "warnings", "ASCII_CITATION_BRACKET", "正文引用默认应使用全角方括号", match.group(0))
    missing = sorted(citation_numbers - reference_numbers)
    if missing:
        add_issue(report, "errors", "CITATION_TARGET_MISSING", "正文引用没有对应的参考文献条目", str(missing))
    record_check(report, "REFERENCES", True, f"检查 {len(blocks)} 条参考文献")


def check_privacy(report: dict, doc) -> None:
    properties = doc.core_properties
    nonempty = {
        key: value
        for key, value in {
            "作者": properties.author,
            "最后修改者": properties.last_modified_by,
            "公司": getattr(properties, "company", ""),
        }.items()
        if value
    }
    if nonempty:
        add_issue(report, "warnings", "DOCUMENT_METADATA", "文档仍含个人或组织元数据，交付前按需清理", str(nonempty))


def audit(path: Path, requested_profile: str, require_baseline: bool = False) -> dict:
    report = {
        "file": str(path),
        "profile": requested_profile,
        "detected_signals": [],
        "checks": [],
        "errors": [],
        "warnings": [],
        "direct_formatting": {},
        "images": {},
        "visual": {
            "status": "视觉待验收",
            "message": "结构审计不包含真实页面渲染，必须另行完成 Word PDF 到 PNG 的视觉验收",
        },
    }
    try:
        with ZipFile(path) as zfile:
            names = zfile.namelist()
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise BadZipFile("不是有效的 Word DOCX 包")
            xml_parts = {
                name: load_xml(zfile, name)
                for name in names
                if name.endswith(".xml") and name.startswith("word/")
            }
            xml_parts["zip_names"] = names
            doc = Document(str(path))
    except (BadZipFile, KeyError, ValueError, OSError) as exc:
        add_issue(report, "errors", "DOCX_OPEN", "无法读取 DOCX", str(exc))
        return report

    detected_profile, signals = detect_profile(doc)
    report["detected_profile"] = detected_profile
    report["detected_signals"] = signals
    if requested_profile == "auto":
        report["profile"] = detected_profile
    elif requested_profile not in {"general", "thesis"}:
        add_issue(report, "errors", "PROFILE", f"未知模式：{requested_profile}")
    else:
        report["profile"] = requested_profile

    check_page_and_sections(report, doc, xml_parts, report["profile"])
    check_baseline_marker(report, xml_parts, require_baseline)
    check_legacy_styles(report, doc)
    check_styles(report, doc, report["profile"])
    check_numbering_bindings(report, doc, xml_parts)
    check_direct_formatting(report, doc)
    check_headings_and_title(report, doc)
    check_appendix(report, doc, xml_parts)
    check_abstracts(report, doc)
    check_toc(report, doc, xml_parts)
    check_captions_tables_images(report, doc, xml_parts)
    check_equations(report, xml_parts)
    check_references_and_citations(report, doc)
    check_privacy(report, doc)

    report["summary"] = {
        "error_count": len(report["errors"]),
        "warning_count": len(report["warnings"]),
        "check_count": len(report["checks"]),
        "result": "结构通过" if not report["errors"] else "发现结构错误",
        "visual_status": report["visual"]["status"],
    }
    return report


def print_report(report: dict) -> None:
    lines = [
        "中文 DOCX 结构审计",
        f"文件：{report['file']}",
        f"模式：{report.get('profile', '未知')}（识别信号：{', '.join(report.get('detected_signals', [])) or '无'}）",
        f"结构结果：{report['summary']['result']}；错误 {report['summary']['error_count']} 项，警告 {report['summary']['warning_count']} 项",
        f"视觉状态：{report['visual']['status']}（{report['visual']['message']}）",
    ]
    if report["errors"]:
        lines.append("\n错误：")
        for item in report["errors"]:
            lines.append(f"- [{item['code']}] {item['message']}" + (f"：{item['details']}" if item.get("details") else ""))
    if report["warnings"]:
        lines.append("\n警告：")
        for item in report["warnings"]:
            lines.append(f"- [{item['code']}] {item['message']}" + (f"：{item['details']}" if item.get("details") else ""))
    output = "\n".join(lines)
    encoding = sys.stdout.encoding or "utf-8"
    print(output.encode(encoding, errors="replace").decode(encoding, errors="replace"))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="只读审计中文 DOCX 的非内容格式")
    parser.add_argument("input", type=Path, help="待审计的 DOCX 文件")
    parser.add_argument("--profile", choices=("auto", "thesis", "general"), default="auto")
    parser.add_argument("--require-baseline", action="store_true", help="要求文档含本技能空白基线模板标识")
    parser.add_argument("--json", dest="json_path", type=Path, help="可选的 JSON 报告路径")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.input.suffix.lower() != ".docx":
        print("输入文件必须是 .docx", file=sys.stderr)
        return 2
    if not args.input.is_file():
        print(f"文件不存在：{args.input}", file=sys.stderr)
        return 2
    report = audit(args.input.resolve(), args.profile, args.require_baseline)
    if args.json_path:
        args.json_path.parent.mkdir(parents=True, exist_ok=True)
        args.json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print_report(report)
    return 1 if report["errors"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
